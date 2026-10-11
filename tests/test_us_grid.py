"""Unit tests for the US grid pack's build steps (no network: small synthetic workbooks)."""

import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from fetch_us_grid import (  # noqa: E402
    build_plants,
    build_queue,
    eia_link,
    fuel_of,
    lbnl_link,
    uswtdb_link,
)
from fetch_us_lines import minzoom  # noqa: E402

PLANT_HEAD = [
    "Entity ID",
    "Plant ID",
    "Plant Name",
    "Plant State",
    "Nameplate Capacity (MW)",
    "Net Summer Capacity (MW)",
    "Technology",
    "Energy Source Code",
    "Status",
    "Operating Year",
    "Planned Retirement Year",
    "Planned Operation Year",
    "Retirement Year",
    "Latitude",
    "Longitude",
]


def sheet(wb: Workbook, name: str, rows: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.append([f"Inventory of {name} Generators as of August 2026"])
    ws.append([])
    ws.append(PLANT_HEAD)
    for r in rows:
        ws.append([r.get(h) if isinstance(r, dict) else None for h in PLANT_HEAD])
    ws.append([])
    ws.append(["NOTES: Capacity from facilities below 1 MW is excluded."])


def unit(pid, name, nameplate, summer, tech, source, status="", **years):
    """One generator row; years by their column (Operating Year=...)."""
    row = {
        "Entity ID": 1,
        "Plant ID": pid,
        "Plant Name": name,
        "Plant State": "TX",
        "Nameplate Capacity (MW)": nameplate,
        "Net Summer Capacity (MW)": summer,
        "Technology": tech,
        "Energy Source Code": source,
        "Status": status,
        "Latitude": 31.0,
        "Longitude": -97.0,
    }
    return {**row, **{k.replace("_", " "): v for k, v in years.items()}}


CC = "Natural Gas Fired Combined Cycle"


def test_eia_link_skips_months_inside_html_comments():
    html = (
        '<!-- <a href="/electricity/data/eia860m/xls/december_generator2026.xlsx">Dec</a> -->'
        '<a href="/electricity/data/eia860m/xls/august_generator2026.xlsx">Aug</a>'
        '<a href="/electricity/data/eia860m/archive/xls/july_generator2026.xlsx">Jul</a>'
    )
    url, release = eia_link(html)
    assert url == "https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx"
    assert release == "August 2026"


def test_lbnl_link_takes_the_newest_year():
    html = (
        '<a href="/sites/default/files/2025-08/lbnl_ix_queue_data_file_thru2024_v2.xlsx">'
        '<a href="/sites/default/files/2026-05/lbnl_ix_queue_data_file_thru2025.xlsx">'
    )
    url, name = lbnl_link(html)
    assert name == "lbnl_ix_queue_data_file_thru2025.xlsx"
    assert url.startswith("https://eta-publications.lbl.gov/sites/default/files/2026-05/")


def test_uswtdb_release_from_the_metadata_file_name():
    assert uswtdb_link('href="../assets/data/uswtdb_V9_1_20260928.xml"')[1] == "V9.1 (2026-09-28)"


def test_fuel_groups():
    assert fuel_of("WAT", "Hydroelectric Pumped Storage") == "storage"
    assert fuel_of("WAT", "Conventional Hydroelectric") == "hydro"
    assert fuel_of("MWH", "Batteries") == "storage"
    assert fuel_of("NG", "Natural Gas Fired Combined Cycle") == "gas"
    assert fuel_of("SUB", "Conventional Steam Coal") == "coal"
    assert fuel_of("PUR", "All Other") == "other"


def test_build_plants_groups_statuses_and_sums_per_plant(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    op = "(OP) Operating"
    sheet(
        wb,
        "Operating",
        [
            unit(
                10, "A", 300, 280, CC, "NG", op, Operating_Year=2001, Planned_Retirement_Year=2030
            ),
            # a combined-cycle part without a net summer value counts 0 MW
            unit(10, "A", 150, " ", CC, "NG", "(SB) Standby/Backup", Operating_Year=1999),
            unit(11, "B", 50, 50, "Solar Photovoltaic", "SUN", op, Operating_Year=2020),
        ],
    )
    sheet(
        wb,
        "Planned",
        [
            unit(
                12,
                "C",
                200,
                200,
                "Batteries",
                "MWH",
                "(U) Under construction",
                Planned_Operation_Year=2027,
            ),
            unit(
                13,
                "D",
                100,
                100,
                "Onshore Wind Turbine",
                "WND",
                "(P) Planned",
                Planned_Operation_Year=2028,
            ),
        ],
    )
    # an early unit without a plant ID or coordinates: in the totals, not on the map
    old = unit(" ", "Old", 58, 58, "Nuclear", "NUC", Retirement_Year=1968)
    old.update({"Latitude": None, "Longitude": None})
    sheet(wb, "Retired", [old])
    path = tmp_path / "august_generator2026.xlsx"
    wb.save(path)
    f = build_plants(path, "August 2026", "test")
    assert f["totals"]["operating"] == {"gas": 280.0, "solar": 50.0}
    assert f["totals"]["construction"] == {"storage": 200.0}
    assert f["totals"]["planned"] == {"wind": 100.0}
    assert f["totals"]["retired"] == {"nuclear": 58.0}
    assert f["status_detail"]["operating"]["SB"] == {"gas": 0.0}
    assert f["additions"] == {
        "2027": {"construction": {"storage": 200.0}},
        "2028": {"planned": {"wind": 100.0}},
    }
    assert f["retirements_planned"] == {"2030": {"gas": 280.0}}
    assert f["left_out"]["units without a net summer capacity"] == 1
    assert f["left_out"]["units without coordinates"] == 1
    # plant A's two gas units are one point: 280 MW, the earliest year, two units
    a = next(p for p in f["points"] if f["plants"][p[6]][0] == "A")
    assert (
        a[2] == f["types"].index("gas")
        and a[3] == 0
        and a[4] == 280.0
        and a[5] == 1999
        and a[7] == 2
    )
    assert len(f["points"]) == 4


def test_build_queue_counts_each_type_of_a_hybrid(tmp_path):
    wb = Workbook()
    intro = wb.active
    intro.title = "Introduction"
    intro.append(["Interconnection Queue Dataset"])
    intro.append(["May 2026"])
    ws = wb.create_sheet("03. Complete Queue Data")
    ws.append(["RETURN TO CONTENTS"])
    head = ["q_status", "state", "IA_phase_clean", "prop_year"]
    head += ["type_1", "type_2", "type_3", "mw_1", "mw_2", "mw_3"]
    ws.append(head)
    rows = [
        ["active", "TX", "IA Executed", 2027, "Solar", "Battery", None, 100, 50, None],
        ["active", "TX", None, 2028, "Gas", None, None, 300, None, None],
        ["suspended", "TX", "Suspended", None, "Wind", None, None, 80, None, None],
        ["withdrawn", "TX", "Withdrawn", None, "Solar", None, None, 999, None, None],
        ["active", None, "Cluster Study", 2027, "Wind", None, None, 40, None, None],
    ]
    for r in rows:
        ws.append(r)
    path = tmp_path / "lbnl_ix_queue_data_file_thru2025.xlsx"
    wb.save(path)
    q = build_queue(path, path.name, "test")
    assert q["edition"] == "May 2026" and q["through"] == "end of 2025"
    by = dict(zip(q["types"], q["total"]["active"], strict=True))
    assert by == {"Gas": 300.0, "Solar": 100.0, "Battery": 50.0, "Wind": 40.0}
    assert q["total"]["generation_mw"] == 440.0 and q["total"]["storage_mw"] == 50.0
    tx = q["states"]["TX"]
    assert tx["active_n"] == 2 and tx["suspended_n"] == 1
    assert dict(zip(q["types"], tx["suspended"], strict=True))["Wind"] == 80.0
    assert dict(zip(q["phases"], tx["phase"], strict=True)) == {
        "Not reported": 300.0,
        "IA Executed": 150.0,
        "Cluster Study": 0.0,
    }
    assert q["no_state"]["active_n"] == 1
    assert q["by_year"]["2027"] == [0.0, 100.0, 50.0, 40.0]


def test_line_zoom_by_voltage():
    assert [minzoom(kv) for kv in (765, 345, 230, 138)] == [0, 0, 3, 5]
