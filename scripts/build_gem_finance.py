"""Who finances coal and gas power: GEM's finance trackers as money flows for the map.

Reads (data/world/gem/, scripts/fetch_gem.py):
  coal-finance      Global Coal Project Finance Tracker: one row per financier, transaction
                    and coal unit; the financier's share per unit (US$), the plant's
                    coordinates
  gas-finance       Gas Finance Tracker: LNG terminals and gas power plants; the financier's
                    share (US$ million); coordinates joined from GEM's LNG terminal tracker
                    (unit ID) and the Integrated Power Tracker (unit ID)

Values are GEM's. Computed here, documented in docs/DATA_SOURCES.md:
  - every amount in US$ million (coal: US$ / 1e6)
  - deals marked "Stopped" are left out; rows without a financier or an amount are left out
  - projects: a plant's or terminal's deals summed (US$ million, number of deals), with its
    three largest financiers
  - flows: international deals (financier's country differs from the project's) summed per
    financier country, project country and fuel; a flow ends at the money-weighted middle
    of the projects it financed in that country (map placement)
  - totals per country: money out (its financiers abroad), money in (from abroad), domestic
  - multilateral / international financiers have no country: counted in projects and lender
    rankings, never drawn as a flow
Country positions for the flows' starts are the middle of each country's largest outline
(map placement), or a fixed point for the few countries without one.

Writes frontend/public/data/eu/gem/finance.json.

    uv run python scripts/build_gem_finance.py
"""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from shapely.geometry import shape

ROOT = Path(__file__).resolve().parents[1]
GEM = ROOT / "data" / "world" / "gem"
OUT = ROOT / "frontend" / "public" / "data" / "eu" / "gem" / "finance.json"
RELEASES = json.loads((ROOT / "scripts" / "gem_releases.json").read_text(encoding="utf-8"))
FUELS = ["coal", "gas"]
# GEM's country names that differ from the map's outlines (ISO alpha-3)
ALIAS = {
    "Russia": "RUS",
    "Vietnam": "VNM",
    "Myanmar": "MMR",
    "Taiwan": "TWN",
    "Czech Republic": "CZE",
    "Czechia": "CZE",
    "Tanzania": "TZA",
    "Guadeloupe": "GLP",
    "DR Congo": "COD",
    "Democratic Republic of the Congo": "COD",
    "Côte d'Ivoire": "CIV",
    "Singapore": "SGP",
    "Hong Kong": "HKG",
    "Mauritius": "MUS",
    "Barbados": "BRB",
    "US": "USA",
    "USA": "USA",
    "United States": "USA",
    "UK": "GBR",
    "United Kingdom": "GBR",
    "Antigua and Barbuda": "ATG",
    "Aruba": "ABW",
    "South Korea": "KOR",
    "Korea": "KOR",
    "Türkiye": "TUR",
    "Turkey": "TUR",
    "Laos": "LAO",
    "Iran": "IRN",
    "Syria": "SYR",
    "Bolivia": "BOL",
    "Venezuela": "VEN",
    "Moldova": "MDA",
    "North Korea": "PRK",
    "Kosovo": "XKX",
    "Bahamas": "BHS",
    "Brunei": "BRN",
    "Macau": "MAC",
}
# countries the simplified outlines lack: a point inside them (map placement)
FIXED = {
    "TWN": [121.0, 23.7],
    "SGP": [103.82, 1.35],
    "HKG": [114.17, 22.32],
    "MUS": [57.55, -20.25],
    "BRB": [-59.55, 13.15],
    "ATG": [-61.8, 17.08],
    "ABW": [-69.97, 12.52],
    "GLP": [-61.55, 16.25],
    "BHS": [-77.4, 24.7],
    "MAC": [113.55, 22.17],
    "BRN": [114.7, 4.5],
}
NONE = {
    "",
    "unknown",
    "not available",
    "n/a",
    "tbd",
    "international",
    "multilateral",
    "corporate",
    "none",
}


def num(v: object) -> float | None:
    try:
        f = float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return f if f == f and f > 0 else None


def year(v: object) -> int | None:
    s = str(v or "")
    for i in range(len(s) - 3):
        if s[i : i + 4].isdigit() and 1950 <= int(s[i : i + 4]) <= 2100:
            return int(s[i : i + 4])
    return None


def clean(v: object) -> str:
    s = str(v or "").strip()
    return "" if s.lower() in NONE else s


class Countries:
    """Country names -> ISO alpha-3 and a map position."""

    def __init__(self) -> None:
        world = json.loads(
            (ROOT / "frontend" / "public" / "data" / "eu" / "world.json").read_text(
                encoding="utf-8"
            )
        )
        self.by_name: dict[str, str] = {}
        self.at: dict[str, list[float]] = dict(FIXED)
        self.names: dict[str, str] = {}
        for f in world["features"]:
            iso = f["properties"]["iso3"]
            self.by_name[f["properties"]["name"].lower()] = iso
            self.names[iso] = f["properties"]["name"]
            geom = shape(f["geometry"])
            biggest = max(getattr(geom, "geoms", [geom]), key=lambda g: g.area)
            p = biggest.representative_point()
            self.at[iso] = [round(p.x, 2), round(p.y, 2)]
        self.unmatched: set[str] = set()

    def iso(self, name: object) -> str | None:
        n = clean(name)
        if not n:
            return None
        iso = ALIAS.get(n) or self.by_name.get(n.lower())
        if not iso:
            self.unmatched.add(n)
            return None
        self.names.setdefault(iso, n)
        return iso


def coal_deals(c: Countries) -> list[dict]:
    path = next((GEM / "coal-finance").glob("*.xlsx"))
    rows = list(
        load_workbook(path, read_only=True, data_only=True)["Data"].iter_rows(values_only=True)
    )
    head_at = next(i for i, r in enumerate(rows) if r and r[0] == "Financier")
    head = [str(h) if h is not None else "" for h in rows[head_at]]
    col = {h: i for i, h in enumerate(head) if h}
    out = []
    for r in rows[head_at + 1 :]:
        r = list(r) + [None] * len(head)
        fin = clean(r[col["Financier"]])
        usd = num(r[col["This Financier's Unit Share per Transaction"]])
        if not fin or usd is None or clean(r[col["Financing Status"]]).lower() == "stopped":
            continue
        out.append(
            {
                "fuel": 0,
                "financier": fin,
                "fin_iso": c.iso(r[col["Financier Country/Area"]]),
                "public": clean(r[col["Financier Public or Private"]]),
                "usd_m": usd / 1e6,
                "year": year(r[col["Close Year"]]),
                "project": clean(r[col["Plant name"]]) or "Coal plant",
                "key": f"coal:{r[col['GEM location ID']]}",
                "proj_iso": c.iso(r[col["Country/Area"]]),
                "status": clean(r[col["Status"]]),
                "lat": r[col["Latitude"]],
                "lon": r[col["Longitude"]],
            }
        )
    return out


def coordinates() -> dict[str, tuple[float, float]]:
    """GEM unit IDs -> (lon, lat): LNG terminal units and Integrated Power Tracker units."""
    at: dict[str, tuple[float, float]] = {}
    lng = next((GEM / "gas-infrastructure").glob("*LNG-Te*.xlsx"))
    rows = load_workbook(lng, read_only=True, data_only=True)["LNG Terminals"].iter_rows(
        values_only=True
    )
    head = list(next(rows))
    for r in rows:
        d = dict(zip(head, r, strict=False))
        try:
            p = (float(d["Longitude"]), float(d["Latitude"]))
        except (TypeError, ValueError):
            continue
        for k in ("UnitID", "ProjectID"):
            if d.get(k):
                at.setdefault(str(d[k]), p)
    gipt = next((GEM / "integrated-power").glob("*.xlsx"))
    rows = load_workbook(gipt, read_only=True, data_only=True)["Power facilities"].iter_rows(
        values_only=True
    )
    head = list(next(rows))
    iu, la, lo = head.index("GEM unit/phase ID"), head.index("Latitude"), head.index("Longitude")
    for r in rows:
        try:
            at.setdefault(str(r[iu]), (float(r[lo]), float(r[la])))
        except (TypeError, ValueError):
            continue
    return at


def gas_deals(c: Countries) -> list[dict]:
    path = next((GEM / "gas-finance").glob("*.xlsx"))
    wb = load_workbook(path, read_only=True, data_only=True)
    at = coordinates()
    out = []
    for sheet, amount, ids in [
        (
            "LNG Terminals",
            "This Financier's Total Share (US$ Million)",
            ("GEM Combo ID", "GEM Terminal ID"),
        ),
        (
            "Gas Power Plants",
            "This Financier's Unit Share Per Transaction (US$ Million)",
            ("GEM Unit ID", "GEM Project ID"),
        ),
    ]:
        rows = list(wb[sheet].iter_rows(values_only=True))
        head = [str(h) if h is not None else "" for h in rows[0]]
        col = {h: i for i, h in enumerate(head) if h}
        # the power-plant sheet names the financier's country "Country" too (its second one)
        countries = [i for i, h in enumerate(head) if h in ("Country", "Financier HQ Country")]
        proj_c, fin_c = countries[0], countries[-1]
        for r in rows[1:]:
            r = list(r) + [None] * len(head)
            fin = clean(r[col["Financier"]])
            usd = num(r[col[amount]])
            if not fin or usd is None or clean(r[col["Finance Status"]]).lower() == "stopped":
                continue
            p = next((at[str(r[col[k]])] for k in ids if str(r[col[k]]) in at), None)
            out.append(
                {
                    "fuel": 1,
                    "financier": fin,
                    "fin_iso": c.iso(r[fin_c]),
                    "public": clean(r[col["Public/Private"]]),
                    "usd_m": usd,
                    "year": year(r[col["Close Year"]]),
                    "project": clean(r[col["Project Name"]]) or "Gas project",
                    "key": f"gas:{r[col[ids[1]]]}",
                    "proj_iso": c.iso(r[proj_c]),
                    "status": clean(r[col["Status"]]),
                    "lat": p[1] if p else None,
                    "lon": p[0] if p else None,
                }
            )
    return out


def main() -> None:
    c = Countries()
    deals = coal_deals(c) + gas_deals(c)
    projects: dict[str, dict] = {}
    flows: dict[tuple, dict] = {}
    lenders: dict[tuple, dict] = {}
    totals: dict[str, dict] = defaultdict(
        lambda: {"out": [0.0, 0.0], "in": [0.0, 0.0], "domestic": [0.0, 0.0]}
    )
    for d in deals:
        try:
            lon, lat = float(d["lon"]), float(d["lat"])
        except (TypeError, ValueError):
            lon = lat = None
        p = projects.setdefault(
            d["key"],
            {
                "fuel": d["fuel"],
                "usd": 0.0,
                "n": 0,
                "name": d["project"],
                "iso": d["proj_iso"],
                "status": d["status"],
                "lon": lon,
                "lat": lat,
                "by": defaultdict(float),
            },
        )
        p["usd"] += d["usd_m"]
        p["n"] += 1
        p["by"][d["financier"]] += d["usd_m"]
        lk = (d["financier"], d["fuel"])
        lender = lenders.setdefault(
            lk, {"iso": d["fin_iso"] or "", "usd": 0.0, "n": 0, "public": d["public"]}
        )
        lender["usd"] += d["usd_m"]
        lender["n"] += 1
        fi, pi = d["fin_iso"], d["proj_iso"]
        if fi and pi and fi == pi:
            totals[pi]["domestic"][d["fuel"]] += d["usd_m"]
        elif fi and pi:
            totals[fi]["out"][d["fuel"]] += d["usd_m"]
            totals[pi]["in"][d["fuel"]] += d["usd_m"]
            f = flows.setdefault(
                (fi, pi, d["fuel"]),
                {"usd": 0.0, "n": 0, "y0": None, "y1": None, "wx": 0.0, "wy": 0.0, "w": 0.0},
            )
            f["usd"] += d["usd_m"]
            f["n"] += 1
            if d["year"]:
                f["y0"] = min(f["y0"] or d["year"], d["year"])
                f["y1"] = max(f["y1"] or d["year"], d["year"])
            if lon is not None:
                f["wx"] += lon * d["usd_m"]
                f["wy"] += lat * d["usd_m"]
                f["w"] += d["usd_m"]
    used = set()
    flow_rows = []
    for (fi, pi, fuel), f in sorted(flows.items(), key=lambda kv: -kv[1]["usd"]):
        end = [round(f["wx"] / f["w"], 2), round(f["wy"] / f["w"], 2)] if f["w"] else c.at.get(pi)
        if not end or fi not in c.at:
            continue
        used.update((fi, pi))
        flow_rows.append([fi, pi, fuel, round(f["usd"], 1), f["n"], f["y0"], f["y1"], end])
    project_rows = []
    for p in sorted(projects.values(), key=lambda p: -p["usd"]):
        if p["lon"] is None:
            continue
        top = sorted(p["by"].items(), key=lambda kv: -kv[1])[:3]
        project_rows.append(
            [
                round(p["lon"], 3),
                round(p["lat"], 3),
                p["fuel"],
                round(p["usd"], 1),
                p["n"],
                p["name"],
                p["iso"] or "",
                p["status"],
                [[n, round(v, 1)] for n, v in top],
            ]
        )
        if p["iso"]:
            used.add(p["iso"])
    lender_rows = sorted(
        (
            [name, v["iso"], fuel, round(v["usd"], 1), v["n"], v["public"]]
            for (name, fuel), v in lenders.items()
        ),
        key=lambda r: -r[3],
    )
    used.update(r[1] for r in lender_rows if r[1])
    used.update(totals)
    data = {
        "source": "Global Energy Monitor, Global Coal Project Finance Tracker, "
        f"{RELEASES.get('coal-finance', {}).get('updated', '')} release, and Gas Finance Tracker, "
        f"{RELEASES.get('gas-finance', {}).get('updated', '')} release (CC BY 4.0)",
        "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
        "fuels": FUELS,
        "unit": "US$ million",
        "countries": {
            iso: {"name": c.names.get(iso, iso), "at": c.at.get(iso)}
            for iso in sorted(used)
            if iso in c.at
        },
        "flows": flow_rows,
        "projects": project_rows,
        "lenders": lender_rows[:400],
        "totals": {
            iso: {k: [round(v[0], 1), round(v[1], 1)] for k, v in t.items()}
            for iso, t in sorted(totals.items())
        },
        "deals": len(deals),
        "unmatched_countries": sorted(c.unmatched),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    coal = sum(d["usd_m"] for d in deals if d["fuel"] == 0) / 1000
    gas = sum(d["usd_m"] for d in deals if d["fuel"] == 1) / 1000
    print(
        f"{OUT.name}: {len(deals):,} deals (coal ${coal:,.0f} bn, gas ${gas:,.0f} bn), {len(flow_rows)} flows, "
        f"{len(project_rows):,} projects, {OUT.stat().st_size / 1e3:.0f} kB; unmatched: {sorted(c.unmatched)}"
    )


if __name__ == "__main__":
    main()
