"""The US grid pack: every power plant, the interconnection queue, every wind turbine, and the
state outlines they are drawn on. Probes: scripts/probe_eia860m.py, probe_lbnl_queue.py,
probe_uswtdb.py (transmission lines: scripts/fetch_us_lines.py, probe_hifld.py).

Sources (all downloads under 20 MB, kept in data/us/, gitignored):
  - EIA-860M, Preliminary Monthly Electric Generator Inventory (public domain): the newest
    month on https://www.eia.gov/electricity/data/eia860m/ (links inside HTML comments are
    placeholders for months not yet out). Generators of 1 MW and more, with coordinates.
  - LBNL "Queued Up" interconnection queue data file (CC BY 4.0, attribution: Lawrence Berkeley
    National Laboratory and GridTracker): the newest lbnl_ix_queue_data_file_thru<YEAR>.xlsx
    on eta-publications.lbl.gov (emp.lbl.gov answers Cloudflare's block page to scripts).
  - USWTDB, U.S. Wind Turbine Database (USGS, LBNL, American Clean Power; public domain):
    the CSV download, release named by its metadata file.
  - U.S. Census Bureau cartographic boundary file, states 1:20M (cb_2024_us_state_20m, KML,
    public domain): the outlines the queue is raised on.

Computed here (documented in docs/DATA_SOURCES.md):
  plants   generators grouped by fuel (EIA energy source code; pumped storage and batteries
           are storage) and status (operating = EIA's Operating sheet: operating, standby,
           out of service; construction = U, V, TS; planned = P, L, T; retired = Retired
           sheet; canceled or postponed left out); summed per plant, fuel and status into one
           point; capacity = net summer capacity (MW), EIA's measure in its published totals;
           sums per state and nationally, planned additions by planned operation year,
           planned retirements by year, retirements by year
  queue    active and suspended requests (q_status), each resource of a hybrid counted under
           its own type (type_1..3 with mw_1..3, as LBNL's "by type" tables do); MW summed
           per state and type, per interconnection phase and per proposed online year
  turbines one row per turbine, compact columns; projects listed once
  states   outlines rounded to 0.001 degrees, a label point inside each

Writes frontend/public/data/us/{plants,queue,turbines,states}.json. Each file records the
release it was built from; a run rebuilds only files whose source has a new release
(--force rebuilds all). In GitHub Actions it writes the rebuilt names to GITHUB_OUTPUT.

    uv run python scripts/fetch_us_grid.py [--only plants queue turbines states] [--force]
    uv run python scripts/fetch_us_grid.py --only plants --eia-file data/us/july_generator2026.xlsx --out /tmp/x
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook
from shapely.geometry import MultiPolygon, Polygon

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "us"
OUT = ROOT / "frontend" / "public" / "data" / "us"
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) Europe-InfraAtlas/0.3 (us grid)"}

EIA_PAGE = "https://www.eia.gov/electricity/data/eia860m/"
LBNL_PAGE = "https://eta-publications.lbl.gov/research-areas/queued-characteristics-power-plants"
USWTDB_PAGE = "https://energy.usgs.gov/uswtdb/data/"
USWTDB_CSV = "https://energy.usgs.gov/uswtdb/assets/data/uswtdbCSV.zip"
CENSUS_STATES = "https://www2.census.gov/geo/tiger/GENZ2024/kml/cb_2024_us_state_20m.zip"

# ------------------------------------------------------------------ EIA-860M

TYPES = [
    "solar",
    "wind",
    "gas",
    "oil",
    "coal",
    "nuclear",
    "hydro",
    "storage",
    "bio",
    "geothermal",
    "other",
]
STATUSES = ["operating", "construction", "planned", "retired"]
# EIA-860 energy source codes -> the map's fuel groups
SOURCE_TYPE = {
    "SUN": "solar",
    "WND": "wind",
    "NG": "gas",
    "OG": "gas",
    "BFG": "gas",
    "PG": "gas",
    "SGP": "gas",
    "DFO": "oil",
    "RFO": "oil",
    "JF": "oil",
    "KER": "oil",
    "PC": "oil",
    "WO": "oil",
    "BIT": "coal",
    "SUB": "coal",
    "LIG": "coal",
    "ANT": "coal",
    "RC": "coal",
    "WC": "coal",
    "SGC": "coal",
    "SC": "coal",
    "NUC": "nuclear",
    "WAT": "hydro",
    "MWH": "storage",
    "WDS": "bio",
    "WDL": "bio",
    "BLQ": "bio",
    "OBS": "bio",
    "OBL": "bio",
    "OBG": "bio",
    "LFG": "bio",
    "AB": "bio",
    "MSW": "bio",
    "SLW": "bio",
    "GEO": "geothermal",
}
# Planned sheet status codes -> construction or planned
PLANNED_STATUS = {
    "U": "construction",
    "V": "construction",
    "TS": "construction",
    "P": "planned",
    "L": "planned",
    "T": "planned",
}
SHEETS = {"Operating": "operating", "Planned": "planned", "Retired": "retired"}
PR_SHEETS = {"Operating_PR": "operating", "Planned_PR": "planned", "Retired_PR": "retired"}


def get(client: httpx.Client, url: str) -> httpx.Response:
    r = client.get(url)
    r.raise_for_status()
    return r


def download(client: httpx.Client, url: str, path: Path) -> Path:
    """Fetch once into data/us/ (the release is in the file name)."""
    if path.exists() and path.stat().st_size > 0:
        print(f"  have {path.name}")
        return path
    path.parent.mkdir(parents=True, exist_ok=True)
    r = get(client, url)
    path.write_bytes(r.content)
    print(f"  downloaded {path.name} ({len(r.content) / 1e6:.1f} MB)")
    return path


def eia_link(html: str) -> tuple[str, str]:
    """The newest month's file on EIA's page (url, "August 2026"); links inside HTML comments
    are placeholders for months not yet published."""
    live = re.sub(r"<!--.*?-->", "", html, flags=re.S)
    link = re.search(r'href="([^"]*/(\w+)_generator(\d{4})\.xlsx)"', live)
    if not link:
        raise SystemExit("EIA-860M: no file link on the page")
    url = str(httpx.URL(EIA_PAGE).join(link.group(1)))
    return url, f"{link.group(2).capitalize()} {link.group(3)}"


def num(v: object) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def year(v: object) -> int | None:
    n = num(v)
    return int(n) if n else None


def fuel_of(source: str, technology: str) -> str:
    if technology == "Hydroelectric Pumped Storage":
        return "storage"
    return SOURCE_TYPE.get(source, "other")


def sheet_rows(wb: Any, name: str) -> list[dict[str, Any]]:
    rows = wb[name].iter_rows(values_only=True)
    head: tuple | None = None
    out = []
    for r in rows:
        if head is None:
            if r and "Plant ID" in r:
                head = r
            continue
        if not r or num(r[head.index("Nameplate Capacity (MW)")]) is None:
            continue  # blank lines and the notes under the table (a few retired units have no plant ID)
        out.append(dict(zip(head, r, strict=False)))
    return out


def add(d: dict, keys: list[str], mw: float) -> None:
    for k in keys[:-1]:
        d = d.setdefault(k, {})
    d[keys[-1]] = d.get(keys[-1], 0.0) + mw


def rounded(d: Any) -> Any:
    if isinstance(d, dict):
        return {k: rounded(v) for k, v in d.items()}
    return round(d, 1) if isinstance(d, float) else d


def build_plants(path: Path, release: str, url: str) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    plants: dict[tuple[int, str, str], dict[str, Any]] = {}
    names: dict[int, int] = {}
    plant_rows: list[list[Any]] = []
    totals: dict = {}
    states: dict = {}
    detail: dict = {}
    additions: dict = {}
    retire_planned: dict = {}
    retired_by_year: dict = {}
    units: dict[str, int] = defaultdict(int)
    left_out = {
        "canceled or postponed units": 0,
        "units without coordinates": 0,
        "units without a net summer capacity": 0,
    }
    for sheet, group in {**SHEETS, **PR_SHEETS}.items():
        if sheet not in wb.sheetnames:
            continue
        for r in sheet_rows(wb, sheet):
            code = str(r.get("Status") or "").strip("( ").split(")")[0]
            status = PLANNED_STATUS.get(code, "planned") if group == "planned" else group
            kind = fuel_of(str(r["Energy Source Code"] or ""), str(r["Technology"] or ""))
            mw = num(r["Net Summer Capacity (MW)"])
            state = str(r["Plant State"] or "")
            if mw is None:
                left_out["units without a net summer capacity"] += 1
                mw = 0.0
            units[status] += 1
            national = not sheet.endswith("_PR")
            if national:
                add(totals, [status, kind], mw)
                add(detail, [status, code or "RE", kind], mw)
            add(states, [state, status, kind], mw)
            if status in ("construction", "planned"):
                y = year(r.get("Planned Operation Year"))
                if national and y:
                    add(additions, [str(y), status, kind], mw)
            elif status == "operating":
                y = year(r.get("Planned Retirement Year"))
                if national and y:
                    add(retire_planned, [str(y), kind], mw)
            elif status == "retired":
                y = year(r.get("Retirement Year"))
                if national and y:
                    add(retired_by_year, [str(y), kind], mw)
            lat, lon = num(r["Latitude"]), num(r["Longitude"])
            if lat is None or lon is None or not isinstance(r["Plant ID"], int):
                left_out["units without coordinates"] += 1
                continue
            pid = r["Plant ID"]
            if pid not in names:
                names[pid] = len(plant_rows)
                plant_rows.append([str(r["Plant Name"] or "").strip(), state, pid])
            p = plants.setdefault(
                (pid, kind, status), {"lon": lon, "lat": lat, "mw": 0.0, "year": None, "units": 0}
            )
            p["mw"] += mw
            p["units"] += 1
            y = year(
                r.get("Retirement Year")
                if status == "retired"
                else r.get("Operating Year") or r.get("Planned Operation Year")
            )
            if y:
                p["year"] = (max if status == "retired" else min)(p["year"] or y, y)
    if "Canceled or Postponed" in wb.sheetnames:
        left_out["canceled or postponed units"] = len(sheet_rows(wb, "Canceled or Postponed"))
    points = [
        [
            round(p["lon"], 4),
            round(p["lat"], 4),
            TYPES.index(kind),
            STATUSES.index(status),
            round(p["mw"], 1),
            p["year"],
            names[pid],
            p["units"],
        ]
        for (pid, kind, status), p in sorted(plants.items(), key=lambda kv: -kv[1]["mw"])
    ]
    return {
        "source": f"U.S. Energy Information Administration, Form EIA-860M, Preliminary Monthly Electric Generator Inventory, {release} (public domain)",
        "release": release,
        "url": url,
        "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
        "measure": "net summer capacity, MW",
        "types": TYPES,
        "statuses": STATUSES,
        # [lon, lat, type, status, MW, year, plant index, units]
        "points": points,
        # [name, state, EIA plant id]
        "plants": plant_rows,
        "units": dict(units),
        "totals": rounded(totals),
        "status_detail": rounded(detail),
        "states": rounded(states),
        "additions": rounded(dict(sorted(additions.items()))),
        "retirements_planned": rounded(dict(sorted(retire_planned.items()))),
        "retired": rounded({y: v for y, v in sorted(retired_by_year.items()) if int(y) >= 2010}),
        "left_out": left_out,
    }


# ------------------------------------------------------------------ LBNL queue

STORAGE = {"Battery", "Other Storage"}


def lbnl_link(html: str) -> tuple[str, str]:
    """The newest queue data file (url, file name): the latest 'thru' year."""
    links = sorted(
        set(re.findall(r'href="([^"]*ix_queue_data_file_thru(\d{4})[^"]*\.xlsx)"', html)),
        key=lambda x: (x[1], x[0]),
    )
    if not links:
        raise SystemExit("LBNL: no queue data file link on the page")
    url = str(httpx.URL(LBNL_PAGE).join(links[-1][0]))
    return url, url.rsplit("/", 1)[1]


def build_queue(path: Path, release: str, url: str) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb["03. Complete Queue Data"].iter_rows(values_only=True)
    head: tuple | None = None
    projects = []
    for r in rows:
        if head is None:
            if r and "q_status" in r:
                head = r
            continue
        if r and any(v is not None for v in r):
            projects.append(dict(zip(head, r, strict=False)))
    intro = " ".join(
        str(v) for r in wb["Introduction"].iter_rows(values_only=True, max_row=40) for v in r if v
    )
    edition = re.search(
        r"(January|February|March|April|May|June|July|August|September|October|November|December) (\d{4})",
        intro,
    )
    through = re.search(r"thru(\d{4})", release)
    mw_by: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    phases: dict[str, float] = defaultdict(float)
    per_state: dict[str, dict[str, Any]] = {}
    no_state = {"active_n": 0, "active": defaultdict(float)}
    by_year: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    counts: dict[str, int] = defaultdict(int)
    for p in projects:
        status = p["q_status"]
        if status not in ("active", "suspended"):
            continue
        counts[status] += 1
        parts = [(p[f"type_{i}"], num(p[f"mw_{i}"]) or 0.0) for i in (1, 2, 3) if p[f"type_{i}"]]
        state = p["state"]
        st = (
            per_state.setdefault(
                state,
                {
                    "active_n": 0,
                    "suspended_n": 0,
                    "active": defaultdict(float),
                    "suspended": defaultdict(float),
                    "phase": defaultdict(float),
                },
            )
            if state
            else None
        )
        if st:
            st[f"{status}_n"] += 1
        elif status == "active":
            no_state["active_n"] += 1
        phase = p["IA_phase_clean"] or "Not reported"
        for kind, mw in parts:
            mw_by[status][kind] += mw
            if st:
                st[status][kind] += mw
            elif status == "active":
                no_state["active"][kind] += mw
            if status == "active":
                phases[phase] += mw
                if st:
                    st["phase"][phase] += mw
                y = p["prop_year"]
                by_year[str(int(y)) if isinstance(y, (int, float)) else "not given"][kind] += mw
    types = [k for k, _ in sorted(mw_by["active"].items(), key=lambda kv: -kv[1])]
    types += [k for k in mw_by["suspended"] if k not in types]
    phase_order = [k for k, _ in sorted(phases.items(), key=lambda kv: -kv[1])]

    def vec(d: dict[str, float], keys: list[str]) -> list[float]:
        return [round(d.get(k, 0.0), 1) for k in keys]

    gen = sum(v for k, v in mw_by["active"].items() if k not in STORAGE)
    sto = sum(v for k, v in mw_by["active"].items() if k in STORAGE)
    return {
        "source": "Lawrence Berkeley National Laboratory and GridTracker (Interconnection.fyi), Queued Up: interconnection queue data file (CC BY 4.0)",
        "release": release,
        "edition": f"{edition.group(1)} {edition.group(2)}" if edition else "",
        "through": f"end of {through.group(1)}" if through else "",
        "url": url,
        "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
        "types": types,
        "storage_types": sorted(STORAGE & set(types)),
        "phases": phase_order,
        "total": {
            "active_n": counts["active"],
            "suspended_n": counts["suspended"],
            "active": vec(mw_by["active"], types),
            "suspended": vec(mw_by["suspended"], types),
            "generation_mw": round(gen, 1),
            "storage_mw": round(sto, 1),
            "phase": vec(phases, phase_order),
        },
        "states": {
            s: {
                "active_n": v["active_n"],
                "suspended_n": v["suspended_n"],
                "active": vec(v["active"], types),
                "suspended": vec(v["suspended"], types),
                "phase": vec(v["phase"], phase_order),
            }
            for s, v in sorted(per_state.items())
        },
        "no_state": {"active_n": no_state["active_n"], "active": vec(no_state["active"], types)},
        "by_year": {y: vec(v, types) for y, v in sorted(by_year.items())},
    }


# ------------------------------------------------------------------ USWTDB


def uswtdb_link(html: str) -> tuple[str, str]:
    """The CSV download and its release, named by the metadata file ("V9.1 (2026-09-28)")."""
    m = re.search(r"uswtdb_V(\d+)_(\d+)_(\d{8})\.xml", html)
    if not m:
        raise SystemExit("USWTDB: no release on the data page")
    d = m.group(3)
    return USWTDB_CSV, f"V{m.group(1)}.{m.group(2)} ({d[:4]}-{d[4:6]}-{d[6:]})"


def build_turbines(path: Path, release: str, url: str) -> dict[str, Any]:
    z = zipfile.ZipFile(path)
    name = next(n for n in z.namelist() if n.endswith(".csv"))
    rows = list(csv.DictReader(io.TextIOWrapper(z.open(name), encoding="utf-8-sig")))

    def val(x: str) -> float | None:
        try:
            v = float(x)
        except ValueError:
            return None
        return None if v < 0 else v  # -9999 marks a missing value in older releases

    projects: dict[tuple[str, str], int] = {}
    project_rows: list[list[Any]] = []
    models: dict[str, int] = {}
    cols: dict[str, list[Any]] = {
        k: [] for k in ("lon", "lat", "kw", "hh", "rd", "project", "model", "offshore")
    }
    states: dict[str, list[float]] = defaultdict(lambda: [0, 0.0])
    for r in rows:
        key = (r["p_name"], r["t_state"])
        if key not in projects:
            projects[key] = len(project_rows)
            pcap = val(r["p_cap"])
            project_rows.append(
                [
                    r["p_name"],
                    r["t_state"],
                    int(val(r["p_year"]) or 0) or None,
                    int(val(r["p_tnum"]) or 0) or None,
                    pcap,
                ]
            )
        model = r["t_model"] or ""
        if model not in models:
            models[model] = len(models)
        kw = val(r["t_cap"])
        cols["lon"].append(round(float(r["xlong"]), 4))
        cols["lat"].append(round(float(r["ylat"]), 4))
        cols["kw"].append(int(kw) if kw is not None else None)
        hh = val(r["t_hh"])
        cols["hh"].append(round(hh, 1) if hh is not None else None)
        rd = val(r["t_rd"])
        cols["rd"].append(round(rd, 1) if rd is not None else None)
        cols["project"].append(projects[key])
        cols["model"].append(models[model])
        cols["offshore"].append(1 if r.get("t_offshore") == "1" else 0)
        s = states[r["t_state"]]
        s[0] += 1
        s[1] += (kw or 0) / 1000
    mw = sum(k for k in cols["kw"] if k) / 1000
    return {
        "source": f"U.S. Wind Turbine Database {release} (USGS, LBNL, American Clean Power; public domain)",
        "release": release,
        "url": url,
        "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
        "count": len(rows),
        "with_capacity": sum(k is not None for k in cols["kw"]),
        "mw": round(mw, 1),
        "offshore_n": sum(cols["offshore"]),
        # one entry per turbine in each column
        **cols,
        # [name, state, year online, turbines in the project, project MW]
        "projects": project_rows,
        "models": list(models),
        # state: [turbines, MW (sum of turbine capacities)]
        "states": {k: [v[0], round(v[1], 1)] for k, v in sorted(states.items())},
    }


# ------------------------------------------------------------------ Census states

KML = "{http://www.opengis.net/kml/2.2}"


def build_states(path: Path) -> dict[str, Any]:
    z = zipfile.ZipFile(path)
    root = ET.fromstring(z.read(next(n for n in z.namelist() if n.endswith(".kml"))))
    out = []
    for pm in root.iter(f"{KML}Placemark"):
        data = {d.get("name"): d.text for d in pm.iter(f"{KML}SimpleData")}
        polys = []
        for poly in pm.iter(f"{KML}Polygon"):
            rings = []
            for tag in ("outerBoundaryIs", "innerBoundaryIs"):
                for ring in poly.iter(f"{KML}{tag}"):
                    coords = (ring.find(f".//{KML}coordinates").text or "").split()
                    rings.append([tuple(float(v) for v in c.split(",")[:2]) for c in coords])
            polys.append(rings)
        shapes = [Polygon(p[0], p[1:]) for p in polys]
        largest = max(shapes, key=lambda s: s.area)
        label = largest.representative_point()
        flat = [[[round(v, 3) for xy in ring for v in xy] for ring in p] for p in polys]
        out.append(
            {
                "code": data["STUSPS"],
                "name": data["NAME"],
                "label": [round(label.x, 3), round(label.y, 3)],
                "bbox": [round(v, 3) for v in MultiPolygon(shapes).bounds],
                "polygons": flat,
            }
        )
    out.sort(key=lambda s: s["code"])
    return {
        "source": "U.S. Census Bureau, cartographic boundary file, states 1:20,000,000 (cb_2024_us_state_20m), public domain",
        "release": "cb_2024_us_state_20m",
        "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
        # polygons: [outer ring, holes...] with flat [lon, lat, lon, lat, ...] rings
        "states": out,
    }


# ------------------------------------------------------------------ main


def current_release(out: Path, name: str) -> str | None:
    f = out / f"{name}.json"
    if not f.exists():
        return None
    try:
        return json.loads(f.read_text(encoding="utf-8")).get("release")
    except json.JSONDecodeError:
        return None


def write(out: Path, name: str, payload: dict[str, Any]) -> None:
    out.mkdir(parents=True, exist_ok=True)
    f = out / f"{name}.json"
    f.write_text(json.dumps(payload, separators=(",", ":"), ensure_ascii=False), encoding="utf-8")
    print(
        f"  {f.relative_to(ROOT) if f.is_relative_to(ROOT) else f}: {f.stat().st_size / 1e6:.2f} MB"
    )


def run_states(client: httpx.Client, args: argparse.Namespace) -> bool:
    if not args.force and current_release(args.out, "states"):
        return False
    print("states (Census)")
    path = download(client, CENSUS_STATES, RAW / CENSUS_STATES.rsplit("/", 1)[1])
    write(args.out, "states", build_states(path))
    return True


def run_plants(client: httpx.Client, args: argparse.Namespace) -> bool:
    if args.eia_file:
        m = re.search(r"(\w+)_generator(\d{4})", args.eia_file.name)
        url = str(args.eia_file)
        release = f"{m.group(1).capitalize()} {m.group(2)}" if m else args.eia_file.stem
    else:
        url, release = eia_link(get(client, EIA_PAGE).text)
    print(f"plants (EIA-860M {release})")
    if not (args.force or args.eia_file or release != current_release(args.out, "plants")):
        print("  release unchanged")
        return False
    path = args.eia_file or download(client, url, RAW / url.rsplit("/", 1)[1])
    payload = build_plants(path, release, url)
    write(args.out, "plants", payload)
    for st in STATUSES:
        gw = sum(payload["totals"].get(st, {}).values()) / 1000
        print(f"  {st:12} {gw:8.1f} GW  (50 states + DC, net summer)")
    print(f"  left out: {payload['left_out']}")
    return True


def run_queue(client: httpx.Client, args: argparse.Namespace) -> bool:
    url, release = lbnl_link(get(client, LBNL_PAGE).text)
    print(f"queue (LBNL {release})")
    if not (args.force or release != current_release(args.out, "queue")):
        print("  release unchanged")
        return False
    payload = build_queue(download(client, url, RAW / release), release, url)
    write(args.out, "queue", payload)
    t = payload["total"]
    print(
        f"  active: {t['active_n']:,} requests, generation {t['generation_mw'] / 1000:,.1f} GW,"
        f" storage {t['storage_mw'] / 1000:,.1f} GW; suspended {t['suspended_n']:,}"
    )
    print(f"  without a state: {payload['no_state']['active_n']} active requests")
    return True


def run_turbines(client: httpx.Client, args: argparse.Namespace) -> bool:
    url, release = uswtdb_link(get(client, USWTDB_PAGE).text)
    print(f"turbines (USWTDB {release})")
    if not (args.force or release != current_release(args.out, "turbines")):
        print("  release unchanged")
        return False
    name = "uswtdb_" + re.sub(r"\W+", "_", release).strip("_") + ".zip"
    payload = build_turbines(download(client, url, RAW / name), release, url)
    write(args.out, "turbines", payload)
    print(
        f"  {payload['count']:,} turbines, {payload['mw'] / 1000:,.1f} GW"
        f" over {payload['with_capacity']:,} with a capacity"
    )
    return True


RUNS = {"states": run_states, "plants": run_plants, "queue": run_queue, "turbines": run_turbines}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--only", nargs="*", choices=list(RUNS), default=list(RUNS))
    parser.add_argument(
        "--force", action="store_true", help="rebuild even if the release did not change"
    )
    parser.add_argument(
        "--eia-file",
        type=Path,
        help="build plants from this EIA-860M file (e.g. an archived month)",
    )
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    changed: list[str] = []
    failed: list[str] = []
    # each source on its own: one that fails (a site down, a changed page) keeps the others
    with httpx.Client(timeout=180, follow_redirects=True, headers=UA) as client:
        for name, run in RUNS.items():
            if name not in args.only:
                continue
            try:
                if run(client, args):
                    changed.append(name)
            except (httpx.HTTPError, SystemExit, KeyError, ValueError) as e:
                print(f"{name} FAILED: {e!r}")
                failed.append(name)
    print(f"changed: {' '.join(changed) or 'nothing'}")
    gh = os.environ.get("GITHUB_OUTPUT")
    if gh:
        with open(gh, "a", encoding="utf-8") as f:
            f.write(f"changed={' '.join(changed)}\n")
    if failed:
        raise SystemExit(f"failed: {' '.join(failed)}")


if __name__ == "__main__":
    main()
