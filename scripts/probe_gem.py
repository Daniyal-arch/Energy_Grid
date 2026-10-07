"""Probe GEM's Global Integrated Power Tracker (September 2026 release) before using it.

Sheets, columns, row counts, and the values of the type and status columns.

    uv run python scripts/probe_gem.py
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "world"
    / "Global Integrated Power September 2026.xlsx"
)
wb = load_workbook(PATH, read_only=True, data_only=True)
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    print(f"== sheet {ws.title!r}: columns {list(header) if header else None}")
    if not header or len(header) < 8:
        for _ in range(3):
            print("   ", next(rows, None))
        continue
    body = list(rows)
    print(f"   rows {len(body)}")
    for col in ("Type", "Status", "Capacity (MW)", "Country/area", "Country/Area"):
        if col in header:
            i = header.index(col)
            vals = Counter(r[i] for r in body)
            print(
                f"   {col}: {len(vals)} distinct, top {vals.most_common(12) if col in ('Type', 'Status') else list(vals.items())[:5]}"
            )
    print("   sample:", body[0])
