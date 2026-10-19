"""What is coming and what it costs (brief 21, group D).

    python scripts/build_coming.py hydrogen|auctions|atb|irena|gas|ghg [--file F] [--lake lake-out] [--upload]

hydrogen  European Hydrogen Observatory (European Commission / Clean Hydrogen Partnership),
          "Hydrogen production and consumption projects" xlsx (the newest linked on its page):
          name, city, country, start year, production capacity (MW), process, end use,
          consumption (t/year), position and stage as published -> frontend/public/data/eu/hydrogen.json
          (layer "Hydrogen projects"), lake `hydrogen/projects`.
auctions  AURES II auction database (Mendeley Data 10.17632/2h9jpzy3wx, CC BY 4.0; EU + UK,
          2011-2021): per auction round the date, technology, status, awarded capacity (kW ->
          MW, computed) and awarded prices in ct/kWh as published -> frontend/public/data/eu/auctions.json,
          lake `aures/auctions` (every column).
atb       NREL Annual Technology Baseline, electricity, newest edition (OEDI, CC BY 4.0): LCOE,
          CAPEX, fixed O&M and capacity factor per technology (default detail), Moderate
          scenario, Market case without tax credits, 30-year cost recovery, for 2025/2030/2035/2050,
          values as published -> frontend/public/data/eu/costs.json ("atb"), lake `atb/electricity`.
irena     IRENA Renewable Power Generation Costs data file (xlsx; irena.org answers scripts with
          403, so scripts/browser_fetch.py irena_costs passes --file): global weighted-average
          LCOE per technology and year as published -> costs.json ("irena").
gas       ENTSOG TYNDP 2026 Annex A (xlsx): every project's code, name, country, promoter,
          maturity, commissioning years and CAPEX (where not confidential) as published; pipeline
          km per project summed over its sections (computed) -> frontend/public/data/eu/gas_projects.json,
          lake `entsog_tyndp/projects`.
ghg       National greenhouse-gas inventories as Eurostat publishes them from the EEA's data
          (`env_air_gge`, API, no key, CC BY 4.0): Mt CO2e by CRF sector, the newest year ->
          frontend/public/data/eu/profiles/ghg_inventory.json (country report). API only.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import wave5_common as w

TMP = Path("/tmp/wave6/coming")
H2_PAGE = "https://observatory.clean-hydrogen.europa.eu/hydrogen-landscape/projects-and-valleys/hydrogen-production-and-consumption-projects"
AURES = "https://data.mendeley.com/public-api/datasets/2h9jpzy3wx/files?folder_id=root&version=1"
ATB = "https://oedi-data-lake.s3.amazonaws.com"
ENTSOG_PAGE = "https://www.entsog.eu/tyndp"
EUROSTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/env_air_gge"
GHG_CODES = {
    "TOTX4_MEMO": "Total, excluding LULUCF",
    "TOTXMEMO": "Total, including LULUCF",
    "CRF1": "Energy",
    "CRF1A1": "of which energy industries (power, heat, refining)",
    "CRF1A2": "of which manufacturing and construction",
    "CRF1A3": "of which transport",
    "CRF1A4": "of which other combustion (buildings, farming)",
    "CRF2": "Industrial processes and product use",
    "CRF3": "Agriculture",
    "CRF4": "Land use, land-use change and forestry (LULUCF)",
    "CRF5": "Waste",
}
ATB_PARAMS = ["LCOE", "CAPEX", "Fixed O&M", "CF"]
ATB_YEARS = [2025, 2030, 2035, 2050]


def iso3(name: object) -> str | None:
    import pycountry

    alias = {"UK": "GBR", "United Kingdom": "GBR", "Czech Republic": "CZE", "Czechia": "CZE", "Netherlands": "NLD", "Turkey": "TUR", "Türkiye": "TUR",
             "Kosovo": "XKX", "Bosnia and Herzegovina": "BIH", "North Macedonia": "MKD", "Moldova": "MDA", "Slovakia": "SVK", "EL": "GRC"}  # fmt: skip
    t = str(name or "").strip()
    if not t:
        return None
    if t in alias:
        return alias[t]
    try:
        return pycountry.countries.lookup(t).alpha_3
    except LookupError:
        return None


def sheet_rows(path: Path, sheet: str | int = 0) -> list[tuple]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    return [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]


def records(rows: list[tuple], first: str) -> list[dict]:
    """Rows below the header row whose first non-empty cell is `first`, as dicts."""
    hi = next(i for i, r in enumerate(rows) if [v for v in r if v not in (None, "")][:1] == [first])
    head = [re.sub(r"\s+", " ", str(v)).strip() if v is not None else None for v in rows[hi]]
    out = []
    for r in rows[hi + 1 :]:
        rec = {h: v for h, v in zip(head, r, strict=False) if h}
        if any(v not in (None, "") for v in rec.values()):
            out.append(rec)
    return out


def to_lake(dataset: str, table: str, rows: list[dict], args: argparse.Namespace, *, source: str, licence: str, sort: list[str],
            version: str = "", lonlat: tuple[str, str] | None = None) -> None:  # fmt: skip
    import duckdb
    import lake

    TMP.mkdir(parents=True, exist_ok=True)
    tmp = TMP / f"{dataset}_{table}.json"
    clean = [
        {k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in r.items()} for r in rows
    ]
    tmp.write_text(json.dumps(clean, ensure_ascii=False, default=str), encoding="utf-8")
    con = duckdb.connect()
    con.execute("INSTALL spatial; LOAD spatial")
    lake.iso_codes(con)
    lk = w.lake_for(dataset, args.lake, source, licence, "scripts/build_coming.py", version)
    geo = (
        f', CASE WHEN TRY_CAST("{lonlat[0]}" AS DOUBLE) IS NOT NULL THEN ST_Point(TRY_CAST("{lonlat[0]}" AS DOUBLE), TRY_CAST("{lonlat[1]}" AS DOUBLE)) END AS geometry'
        if lonlat
        else ""
    )
    lk.write(
        con, table, f"SELECT *{geo} FROM read_json('{tmp}', sample_size = -1, maximum_object_size = 1000000000)",
        sort=sort, geometry="geometry" if lonlat else None,
    )  # fmt: skip
    w.lake_done(lk, args.upload)


# ---------------------------------------------------------------- hydrogen projects


def h2_row(r: dict, procs: list[str], uses: list[str], stages: list[str]) -> list | None:
    """[name, city, ISO3, start year, MW, process, end use, t/year, lon, lat, stage] (indices into the lists)."""
    lon, lat = r.get("Longitude"), r.get("Latitude")
    if not isinstance(lon, int | float) or not isinstance(lat, int | float):
        return None

    def idx(lst: list[str], v: object) -> int | None:
        if v in (None, ""):
            return None
        v = str(v).strip()
        if v not in lst:
            lst.append(v)
        return lst.index(v)

    mw = r.get("Production capacity MW/MWel")
    t = r.get("Consumption capacity (tonnes/year)")
    return [str(r.get("Name") or "").strip(), r.get("City"), iso3(r.get("Country")), r.get("Start year"),
            round(mw, 2) if isinstance(mw, int | float) else None, idx(procs, r.get("Production process")),
            idx(uses, r.get("End-use")), round(t) if isinstance(t, int | float) else None,
            round(lon, 4), round(lat, 4), idx(stages, r.get("Project stage"))]  # fmt: skip


def build_hydrogen(args: argparse.Namespace) -> None:
    import httpx

    html = httpx.get(H2_PAGE, headers=w.UA, timeout=120, follow_redirects=True).text
    links = re.findall(
        r'href="([^"]*[Hh]ydrogen%20production%20and%20consumption%20projects[^"]*\.xlsx)"', html
    )
    if not links:
        raise SystemExit("Observatory: no projects xlsx on the page")
    url = (
        "https://observatory.clean-hydrogen.europa.eu" + links[0]
        if links[0].startswith("/")
        else links[0]
    )
    path = w.download(url, TMP / "h2_projects.xlsx")
    recs = records(sheet_rows(path), "Name")
    procs: list[str] = []
    uses: list[str] = []
    stages: list[str] = []
    rows = [x for x in (h2_row(r, procs, uses, stages) for r in recs) if x]
    print(
        f"  {len(recs)} projects, {len(rows)} with a position; stages {Counter(stages[x[10]] for x in rows if x[10] is not None).most_common()}",
        flush=True,
    )
    w.write_json(w.EU / "hydrogen.json", {
        "fetched": w.now(), "source": "European Hydrogen Observatory, hydrogen production and consumption projects", "url": url,
        "file": Path(url).name.replace("%20", " "), "processes": procs, "end_uses": uses, "stages": stages, "projects": rows,
    })  # fmt: skip
    to_lake("hydrogen", "projects", [{**r, "iso3": iso3(r.get("Country")), "iso2": None, "admin1": None} for r in recs], args,
            source="European Hydrogen Observatory", licence="European Commission re-use policy (Decision 2011/833/EU)", sort=["Name"],
            version=Path(url).name, lonlat=("Longitude", "Latitude"))  # fmt: skip


# ---------------------------------------------------------------- renewable auctions (AURES II)

AURES_COLS = {
    "Country": "country", "Year": "year", "Auction date (deadline) [MM/YEAR]": "date", "Technology focus": "focus",
    "Eligible Technologies": "technology", "Status of the auction": "status",
    "Auctioned product (budget, electricity, installed capacity)": "product",
    "Volume of product [kW if capacity, kWh/a if electricity, EUR if budget]": "volume",
    "Number of submitted bids": "bids", "Average awarded price [ct/ kWh]": "avg_ct_kwh",
    "Lowest awarded price [ct / kWh]": "low_ct_kwh", "Highest awarded price [ct / kWh]": "high_ct_kwh",
    "Awarded capacity [kW]": "awarded_kw", "Auction scheme (incl. main regulation)": "scheme",
}  # fmt: skip


def aures_round(r: dict) -> dict:
    rec = {AURES_COLS[k]: v for k, v in r.items() if k in AURES_COLS}
    d = rec.get("date")
    rec["date"] = (
        d.strftime("%Y-%m")
        if isinstance(d, datetime)
        else (str(d) if d not in (None, "") else None)
    )
    for k in ("volume", "bids", "avg_ct_kwh", "low_ct_kwh", "high_ct_kwh", "awarded_kw"):
        v = rec.get(k)
        rec[k] = v if isinstance(v, int | float) and not isinstance(v, bool) else None
    rec["awarded_mw"] = (
        round(rec["awarded_kw"] / 1000, 3) if rec.get("awarded_kw") is not None else None
    )
    rec["iso3"] = iso3(rec.get("country"))
    return rec


def build_auctions(args: argparse.Namespace) -> None:
    files = w.get_json(AURES)
    f = next(x for x in files if x["filename"].endswith(".xlsx"))  # type: ignore[union-attr]
    path = w.download(f["content_details"]["download_url"], TMP / f["filename"])
    rows = sheet_rows(path, "Database")
    head = [re.sub(r"\s+", " ", str(v)).strip() if v is not None else None for v in rows[0]]
    raw = [{h: v for h, v in zip(head, r, strict=False) if h} for r in rows[1:]]
    raw = [r for r in raw if r.get("Country")]
    rounds = [aures_round(r) for r in raw]
    by: dict[str, list] = defaultdict(list)
    for r in sorted(rounds, key=lambda r: r["date"] or str(r.get("year") or "")):
        if r["iso3"]:
            by[r["iso3"]].append([r["date"] or (str(r["year"]) if r.get("year") else None), r.get("technology"), r.get("status"), r["awarded_mw"],
                                  r["avg_ct_kwh"], r["low_ct_kwh"], r["high_ct_kwh"], r.get("bids")])  # fmt: skip
    print(
        f"  {len(rounds)} rounds, {len(by)} countries; years {min(r.get('year') or 9999 for r in rounds)}-{max(r.get('year') or 0 for r in rounds)}",
        flush=True,
    )
    w.write_json(w.EU / "auctions.json", {
        "fetched": w.now(), "source": "AURES II auction database (Anatolitis et al., Mendeley Data, CC BY 4.0)", "file": f["filename"],
        "doi": "10.17632/2h9jpzy3wx.1",
        # [date YYYY-MM, technologies, status, awarded MW, average, lowest, highest awarded price (ct/kWh), bids]
        "countries": by,
    })  # fmt: skip
    to_lake("aures", "auctions", [{**{k: (str(v) if isinstance(v, datetime) else v) for k, v in r.items()}, "iso3": iso3(r.get("Country")), "iso2": None, "admin1": None} for r in raw],
            args, source="AURES II auction database (Mendeley Data)", licence="CC BY 4.0", sort=["Country", "Year"], version=f["filename"])  # fmt: skip


# ---------------------------------------------------------------- technology costs


def atb_url() -> tuple[str, str]:
    """The newest ATBe.csv on OEDI: (url, edition)."""
    import httpx

    keys: list[str] = []
    for year in range(2030, 2023, -1):
        xml = httpx.get(
            ATB, params={"list-type": "2", "prefix": f"ATB/electricity/csv/{year}/"}, timeout=60
        ).text
        keys = re.findall(r"<Key>([^<]+ATBe\.csv)</Key>", xml)
        if keys:
            key = sorted(keys)[-1]
            return f"{ATB}/{key}", key.split("/csv/")[1].rsplit("/", 1)[0]
    raise SystemExit("ATB: no ATBe.csv listed")


def atb_pick(cases: Counter) -> tuple[str, str]:
    """Market case without tax credits when published, else the R&D case."""
    for case, tax in (("Market", "None"), ("Market", ""), ("R&D", "")):
        if cases.get((case, tax)):
            return case, tax
    raise SystemExit(f"ATB: no usable case in {cases}")


def build_atb(args: argparse.Namespace) -> None:
    import duckdb

    url, edition = atb_url()
    path = w.download(url, TMP / "ATBe.csv")
    con = duckdb.connect()
    con.execute(
        f"CREATE TABLE atb AS SELECT * FROM read_csv('{path}', all_varchar = true, header = true)"
    )
    cases = Counter(
        {
            (c, t or ""): n
            for c, t, n in con.execute(
                "SELECT core_metric_case, tax_credit_case, count(*) FROM atb GROUP BY ALL"
            ).fetchall()
        }
    )
    case, tax = atb_pick(cases)
    params = ", ".join(f"'{p}'" for p in ATB_PARAMS)
    years = ", ".join(f"'{y}'" for y in ATB_YEARS)
    rows = con.execute(
        f"""SELECT display_name, technology, core_metric_parameter, units, core_metric_variable, TRY_CAST(value AS DOUBLE)
            FROM atb WHERE core_metric_case = '{case}' AND coalesce(tax_credit_case, '') = '{tax}' AND crpyears = '30'
              AND scenario = 'Moderate' AND "default" = '1' AND core_metric_parameter IN ({params}) AND core_metric_variable IN ({years})
            ORDER BY 1, 3, 5"""
    ).fetchall()
    techs: dict[str, dict] = defaultdict(
        lambda: {"technology": None, "values": defaultdict(dict), "units": {}}
    )
    for name, tech, par, unit, year, v in rows:
        t = techs[name]
        t["technology"] = tech
        t["units"][par] = unit
        if v is not None:
            t["values"][par][year] = round(v, 4 if par == "CF" else 2)
    print(
        f"  ATB {edition}: case {case}, tax credits {tax or '(none listed)'}; {len(rows)} values, {len(techs)} technologies",
        flush=True,
    )
    out = w.EU / "costs.json"
    old = json.loads(out.read_text(encoding="utf-8")) if out.exists() else {}
    w.write_json(out, {**old, "fetched": w.now(), "atb": {
        "source": f"NREL Annual Technology Baseline {edition} (CC BY 4.0)", "url": url, "edition": edition,
        "case": case, "tax_credits": tax or None, "scenario": "Moderate", "crp_years": 30, "years": ATB_YEARS,
        "technologies": {k: {"technology": v["technology"], "units": v["units"], "values": v["values"]} for k, v in sorted(techs.items())},
    }})  # fmt: skip
    import lake

    lake.iso_codes(con)
    lk = w.lake_for(
        "atb", args.lake, f"NREL ATB {edition}", "CC BY 4.0", "scripts/build_coming.py", edition
    )
    lk.write(
        con, "electricity",
        """SELECT * EXCLUDE (value), TRY_CAST(value AS DOUBLE) AS value, 'US' AS iso2, 'USA' AS iso3, NULL::VARCHAR AS admin1 FROM atb""",
        sort=["technology", "core_metric_parameter", "core_metric_case", "scenario", "core_metric_variable"],
    )  # fmt: skip
    w.lake_done(lk, args.upload)


def irena_table(rows: list[tuple]) -> dict[str, dict[str, float]]:
    """A sheet with a header row of years and technology names in the first text column:
    {technology: {year: value}} for the numeric cells."""
    hi = next(
        (
            i
            for i, r in enumerate(rows)
            if sum(isinstance(v, int) and 2000 <= v <= 2035 for v in r) >= 5
        ),
        None,
    )
    if hi is None:
        return {}
    years = {j: str(v) for j, v in enumerate(rows[hi]) if isinstance(v, int) and 2000 <= v <= 2035}
    out: dict[str, dict[str, float]] = {}
    for r in rows[hi + 1 :]:
        name = next((str(v).strip() for v in r if isinstance(v, str) and v.strip()), None)
        if not name:
            continue
        vals = {
            y: round(float(r[j]), 4)
            for j, y in years.items()
            if j < len(r) and isinstance(r[j], int | float)
        }
        if len(vals) >= 3:
            out[name] = vals
    return out


def build_irena(args: argparse.Namespace) -> None:
    if not args.file:
        raise SystemExit("irena: the file comes from scripts/browser_fetch.py irena_costs (--file)")
    import openpyxl

    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    print(f"  sheets {wb.sheetnames}", flush=True)
    tables = {}
    for ws in wb.worksheets:
        if re.search(r"lcoe", ws.title, re.I) or re.search(r"Fig", ws.title):
            rows = [
                r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)
            ]
            t = irena_table(rows)
            title = next(
                (str(v) for r in rows[:4] for v in r if isinstance(v, str) and len(v) > 12),
                ws.title,
            )
            if t and re.search(r"lcoe|levelised|levelized", title + ws.title, re.I):
                tables[ws.title] = {"title": title, "rows": t}
    if not tables:
        raise SystemExit("irena: no LCOE table found")
    print(f"  {len(tables)} LCOE tables: {list(tables)[:8]}", flush=True)
    out = w.EU / "costs.json"
    old = json.loads(out.read_text(encoding="utf-8")) if out.exists() else {}
    w.write_json(
        out,
        {
            **old,
            "fetched": w.now(),
            "irena": {
                "source": "IRENA Renewable Power Generation Costs (data file)",
                "file": Path(args.file).name,
                "tables": tables,
            },
        },
    )


# ---------------------------------------------------------------- planned gas infrastructure (ENTSOG TYNDP)


def gas_category(code: str) -> str:
    """The project's category from its code (TRA-F-7 -> TRA)."""
    return code.split("-")[0] if code else ""


def build_gas(args: argparse.Namespace) -> None:
    import httpx

    html = httpx.get(ENTSOG_PAGE, headers=w.UA, timeout=120, follow_redirects=True).text
    links = re.findall(r'href="([^"]*Annex%20A%20List%20of%20Projects[^"]*\.xlsx)"', html)
    if not links:
        raise SystemExit("ENTSOG: no Annex A xlsx on the page")
    url = links[0] if links[0].startswith("http") else "https://www.entsog.eu" + links[0]
    path = w.download(url, TMP / "entsog_annex_a.xlsx")
    main = records(sheet_rows(path, "Project Main Info"), "PRJ Code")
    km: dict[str, float] = defaultdict(float)
    for sheet, col in (("TRA", "Length (km)"), ("H2T", "Section Length (km)")):
        for r in records(sheet_rows(path, sheet), "Code"):
            if isinstance(r.get(col), int | float):
                km[r["Code"]] += r[col]
    prj = None
    projects = []
    for r in main:
        prj = r.get("PRJ Code") or prj
        code = r.get("Code")
        if not code:
            continue
        capex = r.get("Project CAPEX (in million €)")
        projects.append([code, r.get("Project Name"), iso3(r.get("Country")), r.get("Promoter"), r.get("Maturity Status"), gas_category(code),
                         r.get("First Incremental Capacity Commissioned"), r.get("Last Incremental Capacity Commissioned"),
                         capex if isinstance(capex, int | float) else None, round(km[code], 1) if code in km else None, prj])  # fmt: skip
    cats = Counter(p[5] for p in projects)
    print(
        f"  {len(projects)} projects; categories {cats.most_common()}; maturity {Counter(p[4] for p in projects).most_common()}",
        flush=True,
    )
    w.write_json(w.EU / "gas_projects.json", {
        "fetched": w.now(), "source": "ENTSOG TYNDP 2026, Annex A list of projects", "url": url, "file": Path(url).name.replace("%20", " "),
        # [code, name, ISO3, promoter, maturity, category, first year, last year, CAPEX M EUR, pipeline km, PRJ group]
        "projects": projects,
        "categories": {"TRA": "gas pipelines and compressors", "LNG": "LNG terminals", "UGS": "underground storage", "H2T": "hydrogen pipelines",
                       "H2L": "hydrogen import terminals", "H2S": "hydrogen storage", "H2E": "electrolysers", "H2M": "hydrogen mobility",
                       "RET": "retrofitting", "BIO": "biomethane", "CO2": "CO2 transport", "SYN": "synthetic methane", "OTH": "other"},
    })  # fmt: skip
    to_lake("entsog_tyndp", "projects", [{**{k: (str(v) if v is not None else None) for k, v in r.items()}, "iso3": iso3(r.get("Country")), "iso2": None, "admin1": None} for r in main if r.get("Code")],
            args, source="ENTSOG TYNDP 2026 Annex A", licence="ENTSOG terms of use (re-use with attribution)", sort=["Code"], version="TYNDP 2026")  # fmt: skip


# ---------------------------------------------------------------- GHG inventories (Eurostat env_air_gge)


def ghg_rows(
    values: dict[tuple[str, ...], float], geo: str, years: list[str]
) -> tuple[str | None, list[list]]:
    """The newest year with a total for a country and its rows [item, Mt CO2e, unit, note]."""
    for y in sorted(years, reverse=True):
        if (geo, "TOTX4_MEMO", y) in values:
            rows = [
                [label, round(values[(geo, code, y)], 3), "Mt CO2e", y]
                for code, label in GHG_CODES.items()
                if (geo, code, y) in values
            ]
            return y, rows
    return None, []


def build_ghg(args: argparse.Namespace) -> None:
    import build_profiles as bp

    codes = "&".join(f"src_crf={c}" for c in GHG_CODES)
    d = w.get_json(f"{EUROSTAT}?unit=MIO_T&airpol=GHG&{codes}&format=JSON&lang=EN", timeout=300)
    raw = bp.jsonstat_values(d)  # type: ignore[arg-type]
    dims = d["id"]  # type: ignore[index]
    gi, si, ti = dims.index("geo"), dims.index("src_crf"), dims.index("time")
    values = {(k[gi], k[si], k[ti]): v for k, v in raw.items()}
    years = sorted({k[2] for k in values})
    geos = sorted({k[0] for k in values})
    out = {}
    newest = Counter()
    for g in geos:
        y, rows = ghg_rows(values, g, years)
        code = (
            iso3(bp.ISO2_EUROSTAT.get(g, g))
            if len(g) == 2
            else ("EU" if g.startswith("EU") else None)
        )
        if y and code:
            out[code] = rows
            newest[y] += 1
    print(f"  {len(out)} countries; newest years {newest.most_common(3)}", flush=True)
    bp.write_profile("ghg_inventory", {
        "title": "Greenhouse-gas inventory", "year": max(newest) if newest else None,
        "source": "National inventories reported to the UNFCCC and the EU (EEA), as Eurostat publishes them (env_air_gge, CC BY 4.0); Mt CO2e (AR5)",
        "rows": out,
    })  # fmt: skip


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("what", choices=["hydrogen", "auctions", "atb", "irena", "gas", "ghg"])
    ap.add_argument("--file", type=Path)
    ap.add_argument("--lake", type=Path, default=Path("lake-out"))
    ap.add_argument("--upload", action="store_true")
    args = ap.parse_args()
    {
        "hydrogen": build_hydrogen,
        "auctions": build_auctions,
        "atb": build_atb,
        "irena": build_irena,
        "gas": build_gas,
        "ghg": build_ghg,
    }[args.what](args)


if __name__ == "__main__":
    main()
