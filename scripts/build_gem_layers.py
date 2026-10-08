"""Map layers from Global Energy Monitor's trackers beyond power plants (CC BY 4.0).

Reads the files scripts/fetch_gem.py keeps in data/world/gem/ (scripts/probe_gem_trackers.py
for their sheets) and writes compact files for the map to frontend/public/data/eu/gem/:

  pipelines_gas.json, pipelines_oil.json   GGIT gas pipelines, GOIT oil/NGL pipelines (routes)
  lng.json            GGIT LNG terminals: units grouped per terminal, type and status
  coal_terminals.json Global Coal Terminals Tracker
  fields.json         GOGET oil and gas fields (field level)
  coal_mines.json     Global Coal Mine Tracker (non-closed mines)
  steel.json          Global Iron and Steel Tracker (plant, status, crude steel capacity)
  cement.json         Global Cement and Concrete Tracker
  chemicals.json      Global Chemicals Inventory
  iron_ore.json       Global Iron Ore Mines Tracker
  methane.json        Global Methane Emitters Tracker: satellite methane plumes

Values are GEM's, unchanged. Computed here, documented in docs/DATA_SOURCES.md:
  - status classes for drawing: operating; construction; planned (proposed, announced,
    discovered, in development, pre-construction); idle (mothballed, idle, shut in).
    Cancelled, shelved, retired, closed, abandoned and decommissioned assets are left out;
    each record keeps GEM's own status for its hover card.
  - LNG: a terminal's units with the same type and status are one point; its capacity is
    their sum (Mtpa).
  - Steel: a plant's capacity rows with the same status are one point (crude steel
    capacity summed, ttpa); the main production equipment is GEM's.
  - Pipeline routes: touching segments joined into lines, then simplified for drawing
    (Douglas-Peucker, 0.01°, about 1 km) and rounded to 0.001°; lengths and capacities
    are GEM's.

    uv run python scripts/build_gem_layers.py
"""

from __future__ import annotations

import json
import re
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from shapely.geometry import MultiLineString, shape
from shapely.geometry.base import BaseGeometry
from shapely.ops import linemerge

ROOT = Path(__file__).resolve().parents[1]
GEM = ROOT / "data" / "world" / "gem"
OUT = ROOT / "frontend" / "public" / "data" / "eu" / "gem"
RELEASES = json.loads((ROOT / "scripts" / "gem_releases.json").read_text(encoding="utf-8"))
CLASSES = ["operating", "construction", "planned", "idle"]
CLASS_OF = {
    "operating": 0,
    "operating pre-retirement": 0,
    "construction": 1,
    "in-development": 2,
    "in development": 2,
    "proposed": 2,
    "announced": 2,
    "discovered": 2,
    "pre-construction": 2,
    "permitted": 2,
    "mothballed": 3,
    "mothballed pre-retirement": 3,
    "idle": 3,
    "idled": 3,
    "shut in": 3,
}
SIMPLIFY_DEG = 0.01


def source(tracker: str) -> str:
    r = RELEASES.get(tracker, {})
    return f"Global Energy Monitor, {r.get('desc', tracker)}, {r.get('updated', '')} release (CC BY 4.0)"


def one(tracker: str, pattern: str) -> Path:
    files = sorted((GEM / tracker).glob(pattern))
    if not files:
        raise SystemExit(f"no {pattern} in {GEM / tracker}: run scripts/fetch_gem.py --pull-r2")
    return files[-1]


def rows(path: Path, sheet: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = [dict(zip(header, r, strict=False)) for r in it if any(v not in (None, "") for v in r)]
    wb.close()
    return out


def num(v: object) -> float | None:
    try:
        f = float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return f if f == f else None  # NaN -> None


def year(v: object) -> int | None:
    m = re.search(r"(18|19|20)\d{2}", str(v or ""))
    return int(m.group()) if m else None


def text(v: object) -> str:
    s = str(v or "").strip()
    return "" if s in {"-", "--", "None", "nan", "n/a", "N/A", "unknown", "Unknown"} else s


def coords(v: object) -> tuple[float, float] | None:
    """GEM's "lat, lon" text -> (lon, lat)."""
    m = re.findall(r"-?\d+(?:\.\d+)?", str(v or ""))
    if len(m) < 2:
        return None
    lat, lon = float(m[0]), float(m[1])
    return (lon, lat) if -90 <= lat <= 90 and -180 <= lon <= 180 else None


def latlon(lat: object, lon: object) -> tuple[float, float] | None:
    la, lo = num(lat), num(lon)
    return (
        (lo, la)
        if la is not None and lo is not None and -90 <= la <= 90 and -180 <= lo <= 180
        else None
    )


class Points:
    """Points of one layer: [lon, lat, kind, class, size, year, name, status, country]."""

    def __init__(self, tracker: str, unit: str, size_label: str) -> None:
        self.tracker, self.unit, self.size_label = tracker, unit, size_label
        self.kinds: list[str] = []
        self.statuses: list[str] = []
        self.points: list[list] = []
        self.left_out = 0

    def add(self, at, kind: str, status: str, size, yr, name: str, country: str) -> None:
        st = text(status).lower()
        cls = CLASS_OF.get(st)
        if at is None or cls is None:
            self.left_out += 1
            return
        kind = text(kind) or "other"
        if kind not in self.kinds:
            self.kinds.append(kind)
        if st not in self.statuses:
            self.statuses.append(st)
        self.points.append(
            [
                round(at[0], 4),
                round(at[1], 4),
                self.kinds.index(kind),
                cls,
                None if size is None else round(size, 3),
                yr,
                text(name),
                self.statuses.index(st),
                text(country),
            ]
        )

    def write(self, name: str) -> None:
        data = {
            "source": source(self.tracker),
            "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
            "classes": CLASSES,
            "kinds": self.kinds,
            "statuses": self.statuses,
            "unit": self.unit,
            "size_label": self.size_label,
            "left_out": self.left_out,
            "points": self.points,
        }
        path = OUT / f"{name}.json"
        path.write_text(
            json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
        )
        print(
            f"{path.name}: {len(self.points):,} points, {self.left_out:,} left out, {path.stat().st_size / 1e3:.0f} kB"
        )


def line_parts(g: BaseGeometry) -> list:
    if g.is_empty:
        return []
    if g.geom_type == "LineString":
        return [g]
    if hasattr(g, "geoms"):
        return [p for sub in g.geoms for p in line_parts(sub)]
    return []


def pipelines(tracker: str, zip_glob: str, out_name: str, cap_key: str, cap_unit: str) -> None:
    zp = one(tracker, zip_glob)
    with zipfile.ZipFile(zp) as z:
        inner = next(n for n in z.namelist() if n.endswith(".geojson") and "__MACOSX" not in n)
        feats = json.loads(z.read(inner))["features"]
    statuses: list[str] = []
    table, paths, left = [], [], 0
    points = 0
    for f in feats:
        p = f["properties"]
        st = text(p.get("Status")).lower()
        if CLASS_OF.get(st) is None or not f.get("geometry"):
            left += 1
            continue
        parts = line_parts(
            linemerge(MultiLineString(line_parts(shape(f["geometry"])))).simplify(
                SIMPLIFY_DEG, preserve_topology=False
            )
        )
        flat_parts = []
        for part in parts:
            flat = [round(c, 3) for xy in part.coords for c in xy[:2]]
            if len(flat) >= 4:
                flat_parts.append(flat)
        if not flat_parts:
            left += 1
            continue
        if st not in statuses:
            statuses.append(st)
        seg = text(p.get("SegmentName"))
        name = text(p.get("PipelineName")) + (
            f" · {seg}" if seg and seg.lower() not in {"main line", "1"} else ""
        )
        cap = num(p.get(cap_key))
        table.append(
            [
                CLASS_OF[st],
                statuses.index(st),
                None if cap is None else round(cap, 2),
                None if (km := num(p.get("LengthMergedKm"))) is None else round(km),
                year(p.get("StartYear1")),
                name,
                text(p.get("CountriesOrAreas")),
            ]
        )
        for flat in flat_parts:
            paths.append([len(table) - 1, *flat])
            points += len(flat) // 2
    data = {
        "source": source(tracker),
        "fetched": datetime.now(UTC).isoformat(timespec="seconds"),
        "classes": CLASSES,
        "statuses": statuses,
        "capacity_unit": cap_unit,
        "left_out": left,
        "rows": table,
        "paths": paths,
    }
    path = OUT / f"{out_name}.json"
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(
        f"{path.name}: {len(table):,} pipelines, {len(paths):,} parts, {points:,} points, {left:,} left out, {path.stat().st_size / 1e6:.2f} MB"
    )


def lng() -> None:
    pts = Points("gas-infrastructure", "Mtpa", "capacity")
    groups: dict[tuple, dict] = {}
    for r in rows(one("gas-infrastructure", "*LNG-Te*.xlsx"), "LNG Terminals"):
        key = (
            r.get("ProjectID") or r.get("TerminalName"),
            text(r.get("FacilityType")),
            text(r.get("Status")).lower(),
        )
        g = groups.setdefault(key, {"r": r, "cap": None, "year": None})
        cap = num(r.get("CapacityinMtpa"))
        if cap is not None:
            g["cap"] = (g["cap"] or 0) + cap
        y = year(r.get("ActualStartYear")) or year(r.get("LatestPlannedStartYear"))
        if y and (g["year"] is None or y < g["year"]):
            g["year"] = y
    for (_, ftype, status), g in groups.items():
        r = g["r"]
        pts.add(
            latlon(r.get("Latitude"), r.get("Longitude")),
            ftype,
            status,
            g["cap"],
            g["year"],
            r.get("TerminalName"),
            r.get("Country/Area"),
        )
    pts.write("lng")


def coal_terminals() -> None:
    pts = Points("coal-terminals", "Mt/y", "capacity")
    for r in rows(one("coal-terminals", "*.xlsx"), "Terminals"):
        pts.add(
            latlon(r.get("Latitude"), r.get("Longitude")),
            r.get("Terminal Type"),
            r.get("Status"),
            num(r.get("Capacity (Mt)")),
            year(r.get("Start Year")),
            r.get("Coal Terminal Name"),
            r.get("Country/Area"),
        )
    pts.write("coal_terminals")


def fields() -> None:
    pts = Points("oil-gas-extraction", "", "")
    for r in rows(one("oil-gas-extraction", "*.xlsx"), "Field-level main data"):
        pts.add(
            latlon(r.get("Latitude"), r.get("Longitude")),
            r.get("Fuel type"),
            r.get("Status"),
            None,
            year(r.get("Production start year")),
            r.get("Unit Name"),
            r.get("Country/Area"),
        )
    pts.write("fields")


def coal_mines() -> None:
    pts = Points("coal-mines", "Mt/y", "capacity")
    for r in rows(one("coal-mines", "Global Coal Mine Tracker*.xlsx"), "Non-closed mines"):
        cap = num(r.get("Capacity (Mtpa)"))
        if cap is None:
            cap = num(r.get("Production (Mtpa)"))
        pts.add(
            latlon(r.get("Latitude"), r.get("Longitude")),
            r.get("Mine Type"),
            r.get("Status"),
            cap,
            year(r.get("Opening Year")),
            r.get("Mine Name"),
            r.get("Country / Area"),
        )
    pts.write("coal_mines")


def steel() -> None:
    plants = {
        r["GEM plant ID"]: r for r in rows(one("iron-steel", "Plant-level*.xlsx"), "Plant data")
    }
    pts = Points("iron-steel", "kt/y", "crude steel capacity")
    groups: dict[tuple, dict] = {}
    for r in rows(one("iron-steel", "Plant-level*.xlsx"), "Plant capacities and status"):
        key = (r.get("GEM plant ID"), text(r.get("Status")).lower())
        g = groups.setdefault(
            key, {"r": r, "cap": None, "kind": text(r.get("Main production equipment"))}
        )
        cap = num(r.get("Nominal crude steel capacity (ttpa)"))
        if cap is not None:
            g["cap"] = (g["cap"] or 0) + cap
    for (pid, status), g in groups.items():
        p = plants.get(pid, {})
        pts.add(
            coords(p.get("Coordinates")),
            g["kind"],
            status,
            g["cap"],
            year(p.get("Start date")),
            g["r"].get("Plant name (English)"),
            g["r"].get("Country/area"),
        )
    pts.write("steel")


def cement() -> None:
    pts = Points("cement", "Mt/y", "cement capacity")
    for r in rows(one("cement", "*.xlsx"), "Final data"):
        pts.add(
            coords(r.get("Coordinates")),
            r.get("Plant type"),
            r.get("Operating status"),
            num(r.get("Cement capacity (million metric tonnes per annum)")),
            year(r.get("Start date")),
            r.get("Plant name (English)"),
            r.get("Country/area"),
        )
    pts.write("cement")


def chemicals() -> None:
    pts = Points("chemicals", "", "")
    for r in rows(one("chemicals", "*.xlsx"), "Plant data"):
        # the inventory lists plants without a status: they count as operating
        pts.add(
            coords(r.get("Coordinates")),
            r.get("Primary products"),
            "operating",
            None,
            None,
            r.get("Plant name (English)"),
            r.get("Country/area"),
        )
    pts.write("chemicals")


def iron_ore() -> None:
    pts = Points("iron-ore-mines", "kt/y", "production 2025")
    for r in rows(one("iron-ore-mines", "Global Iron Ore*.xlsx"), "Main Data"):
        prod = num(r.get("Production 2025 (ttpa)"))
        if prod is None:
            prod = num(r.get("Design capacity (ttpa)"))
        pts.add(
            coords(r.get("Coordinates")),
            "iron ore",
            r.get("Operating status"),
            prod,
            year(r.get("Start date")),
            r.get("Asset name (English)"),
            r.get("Country/area"),
        )
    pts.write("iron_ore")


def methane() -> None:
    pts = Points("methane-emitters", "kg/h", "emission rate")
    for r in rows(one("methane-emitters", "*.xlsx"), "Plumes"):
        # a plume is an observation: drawn as operating, sized by its emission rate
        lat = r.get("Plume Origin Latitude") or r.get("Latitude")
        lon = r.get("Plume Origin Longitude") or r.get("Longitude")
        pts.add(
            latlon(lat, lon),
            r.get("Type of Infrastructure"),
            "operating",
            num(r.get("Emissions (kg/hr)")),
            year(r.get("Observation Date")),
            r.get("GEM Infrastructure Name (Nearby)") or r.get("Name"),
            r.get("Country/Area"),
        )
    pts.write("methane")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    pipelines(
        "gas-infrastructure", "*Gas-Pipelines*.zip", "pipelines_gas", "CapacityBcm/y", "bcm/y"
    )
    pipelines(
        "oil-infrastructure", "*Oil-NGL-Pipelines*.zip", "pipelines_oil", "CapacityBOEd", "boe/d"
    )
    lng()
    coal_terminals()
    fields()
    coal_mines()
    steel()
    cement()
    chemicals()
    iron_ore()
    methane()


if __name__ == "__main__":
    main()
