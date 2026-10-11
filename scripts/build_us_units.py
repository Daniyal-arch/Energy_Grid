"""What each US fossil plant produced, and its CO2, hour by hour through a day.

EPA Clean Air Markets (CAMPD API, https://api.epa.gov/easey/; scripts/probe_epa_campd.py):
  streaming-services/emissions/apportioned/hourly   per unit and hour: gross load (MWh in
      the hour, already multiplied by the operating time), CO2 mass (short tons), operating
      time; `date` + `hour` in the facility's local standard time all year
  streaming-services/facilities/attributes          per unit: latitude/longitude, primary
      fuel, the generators it drives with their nameplate MW
Units are those of EPA's emission-trading programmes (fossil units over 25 MW). Units file a
quarter's data within 30 days after it ends, so the newest complete day is often three months
old; a quarter counts as complete when as many units report its last day as the quarter
before's last day (at least 90 %).

The region of a plant (one of EIA's 13) comes from EIA: CAMPD's facility id is EIA's plant id,
EIA-860M gives the plant's balancing authority, and the EIA-930 reference tables the BA's
region. The time zone comes from the plant's coordinates (tzfpy); its standard offset turns
CAMPD's hours into UTC, placed on the Eastern-Time day of the US replay (build_us_day.py).

Values are passthrough except: a plant's MW per hour = the sum of its units' gross load (a
unit reporting an hour with zero operating time counts 0; a plant with no unit reporting the
hour is null); CO2 in tonnes = short tons x 0.90718474, summed per plant; the plant's
capacity = the sum of the distinct generators' nameplate MW its units list (null if a unit
lists none); its fuel = the primary fuel of its largest unit by nameplate.

Key: sent as the X-Api-Key header only when EPA_API_KEY is set (GitHub Actions, a laptop);
on the cloud the proxy adds it. Never in a URL or a log.

Writes frontend/public/data/eu/usunits/<YYYY-MM-DD>.json (UnitOutputFile, as units/) and
usunits/index.json:
  {"source", "date", "timezone", "start", "step_s", "slots", "fuels",
   "plants": [[lon, lat, fuel, region, capacity MW | null, "name (ST)", units]],
   "output": [[MW | null per hour] per plant], "co2_t": [[t | null per hour] per plant],
   "countries": {region: {"units", "placed"}}, "areas_without_data": []}

    uv run python scripts/build_us_units.py 2026-06-30
    uv run python scripts/build_us_units.py --recent 7    # the newest 7 complete days
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import time
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import httpx
import openpyxl
from dotenv import load_dotenv
from tzfpy import get_tz

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_us_day import STEP, day_window  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "frontend" / "public" / "data" / "eu" / "usunits"
RAW = ROOT / "data" / "us"
API = "https://api.epa.gov/easey"
REFERENCE = "https://www.eia.gov/electricity/930-content/EIA930_Reference_Tables.xlsx"
SHORT_TON_T = 0.90718474
COMPLETE_SHARE = 0.9
SOURCE = (
    "EPA Clean Air Markets Program Data (CAMPD), hourly apportioned emissions and facility attributes; "
    "regions from EIA-860M and the EIA-930 reference tables"
)
FUELS = ["coal", "gas", "oil", "bio", "other"]


def fuel_of(primary: str | None) -> str:
    """CAMPD's primary fuel text to the app's group (the first fuel named counts)."""
    first = (primary or "").split(",")[0].strip().lower()
    if "coal" in first or "coke" in first:
        return "coal"
    if "gas" in first:
        return "gas"
    if "oil" in first:
        return "oil"
    if "wood" in first or "biomass" in first:
        return "bio"
    return "other"


def generators(text: str | None) -> dict[str, float]:
    """'A1ST (191.8), A1CT (170.1)' -> {generator id: nameplate MW}."""
    return {m[1].strip(): float(m[2]) for m in re.finditer(r"([^,()]+)\(([\d.]+)\)", text or "")}


def standard_offset(lon: float, lat: float) -> timedelta:
    """UTC offset of the plant's time zone in January (standard time, as CAMPD reports)."""
    zone = get_tz(lon, lat) or "America/New_York"
    return ZoneInfo(zone).utcoffset(datetime(2026, 1, 15)) or timedelta(0)


class Campd:
    def __init__(self) -> None:
        key = os.environ.get("EPA_API_KEY", "").strip()
        self.client = httpx.Client(timeout=900, headers={"X-Api-Key": key} if key else {})
        self.days: dict[str, list[dict[str, Any]]] = {}

    def get(self, path: str, **params: str) -> httpx.Response:
        """EPA answers 429 to quick successions: wait and ask again."""
        for attempt in range(8):
            r = self.client.get(f"{API}/{path}", params=params)
            if r.status_code != 429:
                time.sleep(1.5)
                return r
            time.sleep(10 * (attempt + 1))
        r.raise_for_status()
        return r

    def hourly(self, day: str) -> list[dict[str, Any]]:
        if day not in self.days:
            r = self.get(
                "streaming-services/emissions/apportioned/hourly", beginDate=day, endDate=day
            )
            r.raise_for_status()
            self.days[day] = r.json()
            print(f"  CAMPD {day}: {len(self.days[day])} unit-hours", flush=True)
        return self.days[day]

    def newest_served(self) -> date:
        """The last day of the newest quarter the API accepts (its 400 answer names it)."""
        far = (datetime.now(UTC).date() + timedelta(days=200)).isoformat()
        r = self.get("streaming-services/emissions/apportioned/hourly", beginDate=far, endDate=far)
        m = re.search(r"quarter ending on (\d\d)/(\d\d)/(\d{4})", r.text)
        if not m:
            raise SystemExit(f"CAMPD: no newest quarter in {r.status_code} {r.text[:200]}")
        return date(int(m[3]), int(m[1]), int(m[2]))


def quarter_end_before(d: date) -> date:
    first = date(d.year, 3 * ((d.month - 1) // 3) + 1, 1)
    return first - timedelta(days=1)


def units_on(rows: list[dict[str, Any]]) -> int:
    return len({(r["facilityId"], r["unitId"]) for r in rows})


def newest_complete(campd: Campd) -> date:
    """The last day of the newest quarter filed as fully as the quarter before."""
    end = campd.newest_served()
    while True:
        before = quarter_end_before(end)
        n, m = units_on(campd.hourly(end.isoformat())), units_on(campd.hourly(before.isoformat()))
        print(f"  units reporting {end}: {n}, {before}: {m}", flush=True)
        if m and n >= COMPLETE_SHARE * m:
            return end
        end = before


def regions(client: httpx.Client) -> dict[int, str]:
    """EIA plant id -> EIA-930 region, via EIA-860M (the file plants.json was built from)."""
    ref = openpyxl.load_workbook(io.BytesIO(client.get(REFERENCE).content), read_only=True)
    ba_region = {
        r[0]: r[3] for r in ref["BAs"].iter_rows(min_row=2, values_only=True) if r[0] and r[3]
    }
    url = json.loads((ROOT / "frontend/public/data/us/plants.json").read_text())["url"]
    path = RAW / url.rsplit("/", 1)[1]
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(client.get(url).content)
    out: dict[int, str] = {}
    wb = openpyxl.load_workbook(path, read_only=True)
    for sheet in ("Operating", "Retired"):
        head = None
        for r in wb[sheet].iter_rows(values_only=True):
            if head is None:
                if r and "Plant ID" in r:
                    head = {name: i for i, name in enumerate(r)}
                continue
            pid, ba = r[head["Plant ID"]], r[head["Balancing Authority Code"]]
            if isinstance(pid, int) and ba in ba_region:
                out.setdefault(pid, ba_region[ba])
    return out


def build(
    day: date, rows: list[dict[str, Any]], attrs: list[dict[str, Any]], region_of: dict[int, str]
) -> dict[str, Any]:
    """One Eastern-Time day of plant output and CO2 from CAMPD unit-hours."""
    start, slots = day_window(day)
    unit_attr = {(a["facilityId"], a["unitId"]): a for a in attrs}
    offsets: dict[int, timedelta] = {}
    mw: dict[int, list[float | None]] = defaultdict(lambda: [None] * slots)
    co2: dict[int, list[float | None]] = defaultdict(lambda: [None] * slots)
    plant_units: dict[int, set[str]] = defaultdict(set)
    # electric units: they drive a generator, or report gross load in some hour (idle hours
    # carry no load, so a unit off all day is known only by its generators)
    electric = {(r["facilityId"], r["unitId"]) for r in rows if r["grossLoad"] is not None}
    electric |= {
        k for k, a in unit_attr.items() if generators(a["associatedGeneratorsAndNameplateCapacity"])
    }
    for r in rows:
        fid, uid = r["facilityId"], r["unitId"]
        a = unit_attr.get((fid, uid))
        if (fid, uid) not in electric or a is None or a["latitude"] is None:
            continue
        if fid not in offsets:
            offsets[fid] = standard_offset(a["longitude"], a["latitude"])
        t = (
            datetime.fromisoformat(r["date"]).replace(tzinfo=UTC)
            + timedelta(hours=r["hour"])
            - offsets[fid]
        )
        k = int((t - start).total_seconds()) // STEP
        if not 0 <= k < slots:
            continue
        plant_units[fid].add(uid)
        running = (r["opTime"] or 0) > 0
        load = (r["grossLoad"] or 0.0) if running else 0.0
        mw[fid][k] = (mw[fid][k] or 0.0) + load
        tons = (r["co2Mass"] or 0.0) if running else 0.0
        co2[fid][k] = (co2[fid][k] or 0.0) + tons * SHORT_TON_T
    plants, output, co2_t = [], [], []
    counts: dict[str, dict[str, int]] = defaultdict(lambda: {"units": 0, "placed": 0})
    for fid in sorted(plant_units):
        units = [unit_attr[(fid, u)] for u in sorted(plant_units[fid])]
        gens: dict[str, float] = {}
        missing = False
        for u in units:
            g = generators(u["associatedGeneratorsAndNameplateCapacity"])
            missing |= not g
            gens.update(g)
        largest = max(
            units,
            key=lambda u: sum(generators(u["associatedGeneratorsAndNameplateCapacity"]).values()),
        )
        region = region_of.get(fid, "")
        counts[region or "none"]["units"] += len(units)
        counts[region or "none"]["placed"] += len(units)
        a = units[0]
        plants.append(
            [
                round(a["longitude"], 4),
                round(a["latitude"], 4),
                FUELS.index(fuel_of(largest["primaryFuelInfo"])),
                region,
                None if missing else round(sum(gens.values()), 1),
                f"{a['facilityName']} ({a['stateCode']})",
                len(units),
            ]
        )
        output.append([None if v is None else round(v, 1) for v in mw[fid]])
        co2_t.append([None if v is None else round(v, 1) for v in co2[fid]])
    return {
        "source": SOURCE,
        "date": day.isoformat(),
        "timezone": "America/New_York",
        "start": start.isoformat(),
        "step_s": STEP,
        "slots": slots,
        "fuels": FUELS,
        "plants": plants,
        "output": output,
        "co2_t": co2_t,
        "countries": dict(counts),
        "areas_without_data": [],
    }


def write_index(out: Path) -> list[str]:
    days = sorted(f.stem for f in out.glob("????-??-??.json"))
    (out / "index.json").write_text(json.dumps({"days": days}), encoding="utf-8")
    return days


def main() -> None:
    load_dotenv()
    parser = argparse.ArgumentParser()
    parser.add_argument("days", nargs="*", help="YYYY-MM-DD (Eastern Time)")
    parser.add_argument(
        "--recent", type=int, help="the newest N complete days; older files are dropped"
    )
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    campd = Campd()
    if args.days:
        days = [date.fromisoformat(d) for d in args.days]
    else:
        newest = newest_complete(campd)
        wanted = [newest - timedelta(days=i) for i in range(args.recent or 2)]
        have = {f.stem for f in args.out.glob("????-??-??.json")}
        days = sorted(d for d in wanted if d.isoformat() not in have)
        print(
            f"newest complete day {newest}; to build: {[d.isoformat() for d in days]}", flush=True
        )
    args.out.mkdir(parents=True, exist_ok=True)
    with httpx.Client(
        timeout=300, follow_redirects=True, headers={"User-Agent": "Europe-InfraAtlas/0.3"}
    ) as client:
        region_of = regions(client) if days else {}
    attrs: dict[int, list[dict[str, Any]]] = {}
    for day in days:
        if day.year not in attrs:
            r = campd.get("streaming-services/facilities/attributes", year=str(day.year))
            r.raise_for_status()
            attrs[day.year] = r.json()
        # an Eastern-Time day reaches back into the previous local day out west
        rows = campd.hourly((day - timedelta(days=1)).isoformat()) + campd.hourly(day.isoformat())
        payload = build(day, rows, attrs[day.year], region_of)
        path = args.out / f"{day.isoformat()}.json"
        path.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
        placed = sum(1 for p in payload["plants"] if p[3])
        print(
            f"{day}: {len(payload['plants'])} plants ({placed} in a region), {path.stat().st_size / 1e3:.0f} kB",
            flush=True,
        )
    if args.recent and not args.days:
        keep = {(newest - timedelta(days=i)).isoformat() for i in range(args.recent)}
        for f in args.out.glob("????-??-??.json"):
            if f.stem not in keep:
                f.unlink()
    print(f"index: {write_index(args.out)}", flush=True)


if __name__ == "__main__":
    main()
