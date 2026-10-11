"""Probe EPA's Clean Air Markets data (CAMPD API) before the US plant-output layer.

1. Hourly apportioned emissions (`streaming-services/emissions/apportioned/hourly`): fields,
   units, the newest date the API serves, and how complete each quarter is (units file
   their data within 30 days after a quarter ends).
2. The clock of `date` + `hour`: on the spring DST day a unit running all day still has
   24 hours (local standard time, no skipped hour).
3. Facility attributes (`streaming-services/facilities/attributes`): latitude/longitude,
   the generators each unit drives and their nameplate MW.
4. Rate limits and paging; size of a day for the whole country.
5. Joining a plant to EIA's 13 regions: CAMPD facility id = EIA plant id (EIA-860M
   "Balancing Authority Code") -> EIA-930 reference tables (BA -> region).

Key: on the cloud the proxy adds X-Api-Key for api.epa.gov; elsewhere EPA_API_KEY is sent
as that header. It is never put in a URL or printed.

    uv run python scripts/probe_epa_campd.py
"""

from __future__ import annotations

import collections
import io
import os
import re
import time
from datetime import date

import httpx
import openpyxl
from dotenv import load_dotenv

load_dotenv()
API = "https://api.epa.gov/easey"
key = os.environ.get("EPA_API_KEY", "").strip()
client = httpx.Client(timeout=600, headers={"X-Api-Key": key} if key else {})


limited = 0


def get(path: str, **params: str) -> httpx.Response:
    """One request; EPA answers 429 to quick successions, so wait and ask again."""
    global limited
    for attempt in range(6):
        r = client.get(f"{API}/{path}", params=params)
        if r.status_code != 429:
            time.sleep(1.5)
            return r
        limited += 1
        time.sleep(10 * (attempt + 1))
    return r


def hourly(day: str, **extra: str) -> tuple[list[dict], httpx.Response, float]:
    t0 = time.monotonic()
    r = get("streaming-services/emissions/apportioned/hourly", beginDate=day, endDate=day, **extra)
    return (r.json() if r.status_code == 200 else []), r, time.monotonic() - t0


print("== 1. Hourly rows")
rows, r, s = hourly("2026-06-30", stateCode="WY")
print(f"   one state, one day: HTTP {r.status_code}, {len(rows)} rows, {s:.1f}s")
print("   fields:", list(rows[0]))
row = next(x for x in rows if x["grossLoad"])
print(
    "   a row:",
    {
        k: row[k]
        for k in (
            "facilityName",
            "facilityId",
            "unitId",
            "date",
            "hour",
            "opTime",
            "grossLoad",
            "co2Mass",
            "co2MassMeasureFlg",
            "primaryFuelInfo",
            "unitType",
        )
    },
)
r = get(
    "streaming-services/emissions/apportioned/hourly", beginDate="2026-10-05", endDate="2026-10-05"
)
print(f"   a date in this quarter: HTTP {r.status_code} {r.text[:240]}")
newest = re.search(r"quarter ending on (\d\d)/(\d\d)/(\d{4})", r.text)
print("   newest date served:", newest and f"{newest[3]}-{newest[1]}-{newest[2]}")
for day in ("2026-03-31", "2026-06-30", "2026-07-31", "2026-08-31", "2026-09-30"):
    rows, r, s = hourly(day)
    units = {(x["facilityId"], x["unitId"]) for x in rows}
    load = sum(x["grossLoad"] or 0 for x in rows)
    print(
        f"   {day}: {len(rows)} rows, {len(units)} units, {len({x['facilityId'] for x in rows})} facilities, gross load {load / 1000:,.0f} GWh, {len(r.content) / 1e6:.0f} MB in {s:.0f}s"
    )
print("   -> Q2 complete; Q3 so far only the units that filed early")

print("\n== 2. Clock: spring DST day 2026-03-08 and the day before")
for day in ("2026-03-07", "2026-03-08"):
    rows, _, _ = hourly(day, stateCode="WY")
    per = collections.defaultdict(set)
    for x in rows:
        if x["opTime"] == 1:
            per[(x["facilityId"], x["unitId"])].add(x["hour"])
    full = [u for u, h in per.items() if len(h) == 24]
    print(
        f"   {day}: units running all 24 hours: {len(full)} of {len(per)}; hours seen {sorted({h for v in per.values() for h in v})}"
    )
print("   -> hour 2 exists on the DST day: hours are local standard time all year")

print("\n== 3. Facility attributes")
r = get("streaming-services/facilities/attributes", year="2026", stateCode="WY")
attrs = r.json()
print(f"   HTTP {r.status_code}, {len(attrs)} units; fields: {list(attrs[0])}")
for a in attrs[:3]:
    print(
        "  ",
        {
            k: a[k]
            for k in (
                "facilityName",
                "facilityId",
                "unitId",
                "latitude",
                "longitude",
                "primaryFuelInfo",
                "maxHourlyHIRate",
                "associatedGeneratorsAndNameplateCapacity",
                "sourceCategory",
            )
        },
    )
r = get("streaming-services/facilities/attributes", year="2026")
every = r.json()
print(
    f"   all states: {len(every)} units, {len({a['facilityId'] for a in every})} facilities, {len(r.content) / 1e6:.1f} MB"
)
print(
    "   no coordinates:",
    sum(1 for a in every if a["latitude"] is None),
    "; no generators listed:",
    sum(1 for a in every if not a["associatedGeneratorsAndNameplateCapacity"]),
)
print(
    "   generator lists, e.g.:",
    [
        a["associatedGeneratorsAndNameplateCapacity"]
        for a in every
        if "," in (a["associatedGeneratorsAndNameplateCapacity"] or "")
    ][:4],
)
print("   primary fuels:", collections.Counter(a["primaryFuelInfo"] for a in every).most_common(14))
print(
    "   source categories:", collections.Counter(a["sourceCategory"] for a in every).most_common(8)
)

print("\n== 4. Rate limits and paging")
codes = []
for _ in range(5):
    r = client.get(
        f"{API}/streaming-services/facilities/attributes",
        params={"year": "2026", "stateCode": "RI"},
    )
    codes.append(r.status_code)
print(
    f"   5 requests back to back: {collections.Counter(codes)}; headers { ({k: v for k, v in r.headers.items() if 'ratelimit' in k.lower()}) }"
)
r = get(
    "emissions-mgmt/emissions/apportioned/hourly",
    beginDate="2026-06-30",
    endDate="2026-06-30",
    page="1",
    perPage="500",
)
print(
    f"   emissions-mgmt (paged): HTTP {r.status_code}, total {r.headers.get('X-Total-Count')} rows, perPage at most 500; streaming-services answer whole"
)

print("\n== 5. Plant -> EIA region")
ref = openpyxl.load_workbook(
    io.BytesIO(
        httpx.get(
            "https://www.eia.gov/electricity/930-content/EIA930_Reference_Tables.xlsx",
            timeout=120,
            follow_redirects=True,
        ).content
    ),
    read_only=True,
)
ba_region = {r[0]: r[3] for r in ref["BAs"].iter_rows(min_row=2, values_only=True) if r[0]}
print(f"   EIA-930 reference: {len(ba_region)} BAs, regions {sorted(set(ba_region.values()))}")
gen = httpx.get(
    "https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx",
    timeout=300,
    follow_redirects=True,
)
wb = openpyxl.load_workbook(io.BytesIO(gen.content), read_only=True)
plant_ba = {}
for r in wb["Operating"].iter_rows(min_row=4, values_only=True):
    if isinstance(r[2], int) and r[8]:
        plant_ba[r[2]] = r[8]
facilities = {a["facilityId"] for a in every}
joined = [f for f in facilities if plant_ba.get(f) in ba_region]
print(
    f"   EIA-860M August 2026 ({len(gen.content) / 1e6:.0f} MB): {len(plant_ba)} plants with a BA; CAMPD facilities placed in a region: {len(joined)} of {len(facilities)}"
)
print("   not joined, e.g.:", sorted(f for f in facilities if f not in joined)[:12])
print(f"\n429 answers waited out above: {limited}; probe run {date.today()}")
