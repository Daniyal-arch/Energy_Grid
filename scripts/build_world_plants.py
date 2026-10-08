"""The world's power plants from Global Energy Monitor, for the World tab.

Source: Global Energy Monitor, Global Integrated Power Tracker, newest release
(CC BY 4.0): the Excel file scripts/fetch_gem.py keeps in data/world/gem/integrated-power/
(it goes through GEM's download form). 183,404 units in September 2026 (scripts/probe_gem.py).

Kept (computed here, documented in docs/DATA_SOURCES.md):
  - statuses grouped: operating; construction; planned = pre-construction + announced;
    retired = retired + mothballed. Cancelled and shelved projects are left out.
  - units with coordinates, summed per GEM location, type and status group: one point
    each (capacity in MW = the sum of its units; first year = earliest start year;
    for retired, the latest retired year)
  - per country (ISO alpha-3, GEM's country table) and world: capacity by type and status

Writes frontend/public/data/eu/:
  world_plants.json  {"source", "types", "statuses", "points": [[lon, lat, type, status, MW, year]],
                      "countries": {ISO3: {status: {type: MW}}}, "world": {status: {type: MW}}}
  world_plant_names.json  {"names": [plant name per point, same order]}

    uv run python scripts/build_world_plants.py
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
# the newest release fetched by scripts/fetch_gem.py (one Excel file in the folder)
GEM_DIR = ROOT / "data" / "world" / "gem" / "integrated-power"
OUT = ROOT / "frontend" / "public" / "data" / "eu"


def source_file() -> Path:
    files = sorted(GEM_DIR.glob("*.xlsx"))
    if not files:
        raise SystemExit(f"no Excel file in {GEM_DIR}: run scripts/fetch_gem.py first")
    return files[-1]


def source_label(path: Path) -> str:
    """The release as GEM names its file, e.g. "September 2026"."""
    release = path.stem.replace("Global Integrated Power", "").strip() or path.stem
    return f"Global Energy Monitor, Global Integrated Power Tracker, {release} release (CC BY 4.0)"


TYPES = {
    "utility-scale solar": "solar",
    "wind": "wind",
    "oil/gas": "gas",
    "coal": "coal",
    "hydropower": "hydro",
    "bioenergy": "bio",
    "nuclear": "nuclear",
    "geothermal": "geothermal",
}
STATUS = {
    "operating": "operating",
    "construction": "construction",
    "pre-construction": "planned",
    "announced": "planned",
    "retired": "retired",
    "mothballed": "retired",
}
TYPE_ORDER = list(dict.fromkeys(TYPES.values()))
STATUS_ORDER = ["operating", "construction", "planned", "retired"]


def main() -> None:
    src = source_file()
    wb = load_workbook(src, read_only=True, data_only=True)
    countries = wb["Regions, area, and countries"].iter_rows(values_only=True)
    head = next(countries)
    name_i, iso_i = head.index("GEM Standard Country Name/Area"), head.index("ISO-alpha3 Code")
    iso3 = {r[name_i]: r[iso_i] for r in countries if len(r) > iso_i and r[name_i] and r[iso_i]}

    rows = wb["Power facilities"].iter_rows(values_only=True)
    head = next(rows)
    col = {name: head.index(name) for name in head if name}
    plants: dict[tuple, dict] = {}
    per_country: dict[str, dict[str, dict[str, float]]] = {}
    world: dict[str, dict[str, float]] = {}
    skipped = {"status": 0, "no coordinates": 0, "no capacity": 0}
    for r in rows:
        kind = TYPES.get(r[col["Type"]])
        status = STATUS.get(r[col["Status"]])
        if not kind or not status:
            skipped["status"] += 1
            continue
        mw = r[col["Capacity (MW)"]]
        if not isinstance(mw, (int, float)) or mw <= 0:
            skipped["no capacity"] += 1
            continue
        code = iso3.get(r[col["Country/area"]])
        if code:
            cell = per_country.setdefault(code, {}).setdefault(status, {})
            cell[kind] = cell.get(kind, 0.0) + mw
        w = world.setdefault(status, {})
        w[kind] = w.get(kind, 0.0) + mw
        lat, lon = r[col["Latitude"]], r[col["Longitude"]]
        if not isinstance(lat, (int, float)) or not isinstance(lon, (int, float)):
            skipped["no coordinates"] += 1
            continue
        key = (r[col["GEM location ID"]] or f"{lat:.4f},{lon:.4f}", kind, status)
        p = plants.setdefault(
            key,
            {
                "lon": lon,
                "lat": lat,
                "mw": 0.0,
                "year": None,
                "name": r[col["Plant / Project name"]] or "",
            },
        )
        p["mw"] += mw
        year = r[col["Retired year"]] if status == "retired" else r[col["Start year"]]
        if isinstance(year, (int, float)):
            y = int(year)
            p["year"] = max(p["year"] or y, y) if status == "retired" else min(p["year"] or y, y)
    points = []
    names = []
    for (_, kind, status), p in sorted(plants.items(), key=lambda kv: -kv[1]["mw"]):
        points.append(
            [
                round(p["lon"], 4),
                round(p["lat"], 4),
                TYPE_ORDER.index(kind),
                STATUS_ORDER.index(status),
                round(p["mw"], 1),
                p["year"],
            ]
        )
        names.append(p["name"])

    def rounded(d: dict[str, dict[str, float]]) -> dict:
        return {
            s: {k: round(v, 1) for k, v in sorted(t.items(), key=lambda kv: -kv[1])}
            for s, t in d.items()
        }

    fetched = datetime.now(UTC).isoformat(timespec="seconds")
    payload = {
        "source": source_label(src),
        "fetched": fetched,
        "types": TYPE_ORDER,
        "statuses": STATUS_ORDER,
        "points": points,
        "countries": {c: rounded(v) for c, v in sorted(per_country.items())},
        "world": rounded(world),
    }
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "world_plants.json").write_text(
        json.dumps(payload, separators=(",", ":")), encoding="utf-8"
    )
    (OUT / "world_plant_names.json").write_text(
        json.dumps(
            {"source": source_label(src), "fetched": fetched, "names": names},
            separators=(",", ":"),
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    size = (OUT / "world_plants.json").stat().st_size / 1e6
    print(f"world_plants.json: {len(points):,} points, {size:.1f} MB; skipped {skipped}")
    for s in STATUS_ORDER:
        print(
            f"  {s:12} {sum(world.get(s, {}).values()) / 1000:8.0f} GW  {dict(list(rounded(world).get(s, {}).items())[:4])}"
        )


if __name__ == "__main__":
    main()
