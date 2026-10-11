"""Probe EIA-860M (Preliminary Monthly Electric Generator Inventory) before the US plants layer.

1. The newest file on https://www.eia.gov/electricity/data/eia860m/ (links inside HTML
   comments are placeholders for months not yet published: skipped)
2. Its sheets, header row, columns and row counts
3. Status codes, technologies and energy sources used; plants with coordinates
4. Capacity totals per sheet (nameplate and net summer), for the check against EIA's
   published national total

  uv run python scripts/probe_eia860m.py [--file path.xlsx]
"""

from __future__ import annotations

import argparse
import re
import tempfile
from collections import Counter
from pathlib import Path

import httpx
from openpyxl import load_workbook

PAGE = "https://www.eia.gov/electricity/data/eia860m/"


def newest_link(html: str) -> str:
    live = re.sub(r"<!--.*?-->", "", html, flags=re.S)
    links = re.findall(r'href="([^"]*_generator\d{4}\.xlsx)"', live)
    print(f"links outside comments: {len(links)}; first three: {links[:3]}")
    return httpx.URL(PAGE).join(links[0]).__str__()


def header_row(rows: list[tuple]) -> int:
    for i, r in enumerate(rows[:10]):
        if r and "Plant ID" in r:
            return i
    raise SystemExit("no header row with 'Plant ID'")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", type=Path)
    args = parser.parse_args()
    if args.file:
        path = args.file
    else:
        with httpx.Client(timeout=120, follow_redirects=True) as c:
            url = newest_link(c.get(PAGE).text)
            print(f"newest: {url}")
            r = c.get(url)
            r.raise_for_status()
            path = Path(tempfile.gettempdir()) / url.rsplit("/", 1)[1]
            path.write_bytes(r.content)
            print(
                f"  {len(r.content) / 1e6:.1f} MB, Last-Modified {r.headers.get('last-modified')}"
            )
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"sheets: {wb.sheetnames}")
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        if not rows:
            continue
        try:
            h = header_row(rows)
        except SystemExit:
            print(f"\n[{name}] {len(rows)} rows, no generator table; first rows: {rows[:3]}")
            continue
        head = rows[h]
        data = [r for r in rows[h + 1 :] if r and r[head.index("Plant ID")] not in (None, "")]
        print(f"\n[{name}] header on row {h + 1}; {len(data):,} generator rows")
        print(f"  title rows: {[r[0] for r in rows[:h] if r and r[0]]}")
        print(f"  columns: {list(head)}")
        col = {c: i for i, c in enumerate(head) if c}
        for c in (
            "Status",
            "Technology",
            "Energy Source Code",
            "Prime Mover Code",
            "Balancing Authority Code",
        ):
            if c in col:
                print(f"  {c}: {Counter(r[col[c]] for r in data).most_common(40)}")
        for cap in (
            "Nameplate Capacity (MW)",
            "Net Summer Capacity (MW)",
            "Net Winter Capacity (MW)",
        ):
            if cap in col:
                vals = [r[col[cap]] for r in data]
                num = [v for v in vals if isinstance(v, (int, float))]
                print(
                    f"  {cap}: sum {sum(num) / 1000:,.1f} GW, {len(vals) - len(num)} non-numeric {Counter(v for v in vals if not isinstance(v, (int, float))).most_common(5)}"
                )
        if "Latitude" in col:
            ok = sum(
                isinstance(r[col["Latitude"]], (int, float))
                and isinstance(r[col["Longitude"]], (int, float))
                for r in data
            )
            plants = {r[col["Plant ID"]] for r in data}
            print(f"  rows with coordinates: {ok:,} of {len(data):,}; plants: {len(plants):,}")
        if "Plant State" in col:
            print(
                f"  states: {len({r[col['Plant State']] for r in data})}, e.g. {Counter(r[col['Plant State']] for r in data).most_common(6)}"
            )
        for c in (
            "Planned Operation Year",
            "Planned Retirement Year",
            "Retirement Year",
            "Operating Year",
        ):
            if c in col:
                print(
                    f"  {c}: {sorted(Counter(r[col[c]] for r in data if r[col[c]]).items(), key=lambda kv: str(kv[0]))[-12:]}"
                )
        print(f"  first row: {dict(zip(head, data[0], strict=False))}")


if __name__ == "__main__":
    main()
