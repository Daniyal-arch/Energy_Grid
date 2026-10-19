"""US plant performance and emitters (brief 17, group B).

    python scripts/build_us_perf.py perf  [--lake lake-out] [--upload]
    python scripts/build_us_perf.py ghgrp [--lake lake-out] [--upload]

perf   EIA-923 through PUDL's public Parquet (Catalyst Cooperative, CC BY 4.0), read in place with
       DuckDB: per EIA plant, the newest full year's net generation (MWh, the sum of
       `core_eia923__monthly_generation_fuel`) and the capacity factor: PUDL's own
       (`out_eia__yearly_generators.capacity_factor`, capacity-weighted over the plant's generators)
       where every generator has one, else computed: generation / (nameplate MW x hours of the
       year), labelled. Plus EPA eGRID (the newest plant file, xlsx): the plant's annual CO2 (short
       tons) and CO2 output rate (lb/MWh) as published. -> frontend/public/data/us/plants_perf.json;
       lake `egrid/plants` (every column of the PLNT sheet, eGRID's codes as names).
ghgrp  EPA GHGRP: facilities and positions from Envirofacts `PUB_DIM_FACILITY` (the newest year),
       emissions from the data summary spreadsheet (`Direct Point Emitters`, t CO2e, AR4 GWP) ->
       frontend/public/data/us/ghgrp.json (the layer "US industrial emitters"); lake
       `ghgrp/facilities`.
"""

from __future__ import annotations

import argparse
import io
import json
import zipfile
from pathlib import Path

import wave5_common as w

PUDL = "https://s3.us-west-2.amazonaws.com/pudl.catalyst.coop/stable"
EGRID_PAGE = "https://www.epa.gov/egrid/download-data"
EGRID_URL = "https://www.epa.gov/system/files/documents/2025-06/egrid2023_data_rev2.xlsx"
GHGRP_URL = (
    "https://www.epa.gov/system/files/other-files/2024-10/2023_data_summary_spreadsheets.zip"
)
GHGRP_YEAR = 2023
ENVIROFACTS = "https://data.epa.gov/efservice/PUB_DIM_FACILITY/YEAR/=/{year}/ROWS/{a}:{b}/JSON"
TMP = Path("/tmp/wave5")
# GHGRP "Industry Type (sectors)": the first sector a facility lists decides its colour
SECTORS = [
    "Power Plants",
    "Petroleum and Natural Gas Systems",
    "Refineries",
    "Chemicals",
    "Minerals",
    "Metals",
    "Waste",
    "Pulp and Paper",
    "Other",
]


def hours_in(year: int) -> int:
    return 8784 if (year % 4 == 0 and year % 100 != 0) or year % 400 == 0 else 8760


def capacity_factor(gen_mwh: float | None, mw: float | None, year: int) -> float | None:
    """Generation / (nameplate x hours of the year), 3 decimals; None without both."""
    if gen_mwh is None or not mw:
        return None
    return round(gen_mwh / (mw * hours_in(year)), 3)


def sector_of(sectors: str | None) -> int:
    first = (sectors or "").split(",")[0].strip()
    return SECTORS.index(first) if first in SECTORS else SECTORS.index("Other")


def egrid_rows(sheet_rows: list[tuple]) -> list[dict]:
    """eGRID's PLNT sheet: row 0 long names, row 1 codes, then one row per plant."""
    codes = [str(c) if c is not None else f"col{i}" for i, c in enumerate(sheet_rows[1])]
    return [dict(zip(codes, r, strict=False)) for r in sheet_rows[2:] if r and r[0] is not None]


def perf(args: argparse.Namespace) -> None:
    import duckdb
    import openpyxl

    con = duckdb.connect()
    con.execute("INSTALL httpfs; LOAD httpfs")
    gf = f"'{PUDL}/core_eia923__monthly_generation_fuel.parquet'"
    year = con.execute(
        f"""SELECT max(y) FROM (SELECT year(report_date) y, count(DISTINCT month(report_date)) m
            FROM {gf} GROUP BY 1) WHERE m = 12"""
    ).fetchone()[0]
    print(f"  newest full year in EIA-923: {year}", flush=True)
    con.execute(
        f"""CREATE TABLE gen AS SELECT plant_id_eia, sum(net_generation_mwh) AS mwh
            FROM {gf} WHERE year(report_date) = {year} GROUP BY 1"""
    )
    con.execute(
        f"""CREATE TABLE cap AS SELECT plant_id_eia, sum(capacity_mw) AS mw,
                   CASE WHEN bool_and(capacity_factor IS NOT NULL) AND sum(capacity_mw) > 0
                        THEN sum(capacity_factor * capacity_mw) / sum(capacity_mw) END AS cf_pudl
            FROM '{PUDL}/out_eia__yearly_generators.parquet'
            WHERE year(report_date) = {year} AND capacity_mw > 0
              AND (generator_retirement_date IS NULL OR year(generator_retirement_date) >= {year})
            GROUP BY 1"""
    )
    rows = con.execute(
        "SELECT g.plant_id_eia, g.mwh, c.mw, c.cf_pudl FROM gen g LEFT JOIN cap c USING (plant_id_eia)"
    ).fetchall()
    src = w.download(EGRID_URL, TMP / "egrid.xlsx")
    wb = openpyxl.load_workbook(src, read_only=True)
    sheet = next(n for n in wb.sheetnames if n.startswith("PLNT"))
    eg = egrid_rows(list(wb[sheet].iter_rows(values_only=True)))
    egrid_year = int(eg[0]["YEAR"])
    by_id = {int(r["ORISPL"]): r for r in eg if r.get("ORISPL") is not None}
    plants: dict[str, list] = {}
    for pid, mwh, mw, cf_pudl in rows:
        cf = round(cf_pudl, 3) if cf_pudl is not None else capacity_factor(mwh, mw, year)
        e = by_id.get(int(pid), {})
        plants[str(pid)] = [
            round(mwh) if mwh is not None else None,
            cf,
            1 if cf_pudl is not None else 0,
            e.get("PLCO2RTA"),
            round(e["PLCO2AN"]) if e.get("PLCO2AN") is not None else None,
        ]
    for pid, e in by_id.items():  # plants in eGRID without EIA-923 rows that year
        plants.setdefault(
            str(pid),
            [
                None,
                None,
                0,
                e.get("PLCO2RTA"),
                round(e["PLCO2AN"]) if e.get("PLCO2AN") is not None else None,
            ],
        )
    w.write_json(
        w.US / "plants_perf.json",
        {
            "fetched": w.now(),
            "year": year,
            "egrid_year": egrid_year,
            "source": f"EIA-923 via PUDL (Catalyst Cooperative, CC BY 4.0), {year}; EPA eGRID{egrid_year}",
            "fields": [
                "net generation MWh",
                "capacity factor",
                "1 = PUDL's factor, 0 = computed",
                "CO2 lb/MWh (eGRID)",
                "CO2 short tons (eGRID)",
            ],
            "plants": plants,
        },
    )
    print(
        f"  {len(rows):,} plants with EIA-923 generation, {len(by_id):,} in eGRID{egrid_year}",
        flush=True,
    )
    tmp = TMP / "egrid_plants.json"
    tmp.write_text(json.dumps(eg, default=str), encoding="utf-8")
    con.execute("INSTALL spatial; LOAD spatial")
    lk = w.lake_for(
        "egrid",
        args.lake,
        "EPA eGRID",
        "public domain (US EPA)",
        "scripts/build_us_perf.py",
        f"eGRID{egrid_year}",
    )
    lk.write(
        con,
        "plants",
        f"""SELECT *, 'US' AS iso2, 'USA' AS iso3, PSTATABB AS admin1,
                   CASE WHEN TRY_CAST(LON AS DOUBLE) IS NOT NULL THEN ST_Point(TRY_CAST(LON AS DOUBLE), TRY_CAST(LAT AS DOUBLE)) END AS geometry
            FROM read_json('{tmp}', maximum_object_size = 1000000000, sample_size = -1)""",
        sort=["ORISPL"],
        partition={"year": str(egrid_year)},
        geometry="geometry",
    )
    w.lake_done(lk, args.upload)


def ghgrp(args: argparse.Namespace) -> None:
    import duckdb
    import openpyxl

    facilities: dict[int, dict] = {}
    for a in range(0, 30000, 5000):
        part = w.get_json(ENVIROFACTS.format(year=GHGRP_YEAR, a=a, b=a + 4999), timeout=300)
        for r in part:  # type: ignore[union-attr]
            facilities[int(r["facility_id"])] = r
        print(f"  Envirofacts rows {a}..: {len(part)}", flush=True)  # type: ignore[arg-type]
        if len(part) < 5000:  # type: ignore[arg-type]
            break
    src = w.download(GHGRP_URL, TMP / "ghgrp.zip")
    with zipfile.ZipFile(src) as z:
        wb = openpyxl.load_workbook(
            io.BytesIO(z.read(f"ghgp_data_{GHGRP_YEAR}.xlsx")), read_only=True
        )
    sheet = list(wb["Direct Point Emitters"].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(sheet) if r and r[0] == "Facility Id")
    cols = [str(c).strip() for c in sheet[head]]
    rows = [dict(zip(cols, r, strict=False)) for r in sheet[head + 1 :] if r and r[0] is not None]
    points, out = [], []
    for r in rows:
        fid = int(r["Facility Id"])
        f = facilities.get(fid, {})
        lon = f.get("longitude") if f.get("longitude") is not None else r.get("Longitude")
        lat = f.get("latitude") if f.get("latitude") is not None else r.get("Latitude")
        t = r.get("Total reported direct emissions")
        rec = {
            "facility_id": fid,
            "name": r.get("Facility Name"),
            "state": r.get("State"),
            "naics": str(r.get("Primary NAICS Code") or ""),
            "sectors": r.get("Industry Type (sectors)"),
            "subparts": r.get("Industry Type (subparts)"),
            "t_co2e": t,
            "t_co2": r.get("CO2 emissions (non-biogenic)"),
            "t_ch4_co2e": r.get("Methane (CH4) emissions"),
            "t_n2o_co2e": r.get("Nitrous Oxide (N2O) emissions"),
            "lon": lon,
            "lat": lat,
            "position": "Envirofacts" if f.get("longitude") is not None else "data summary",
        }
        out.append(rec)
        if lon is not None and lat is not None and t:
            points.append(
                [
                    round(float(lon), 4),
                    round(float(lat), 4),
                    round(float(t)),
                    sector_of(rec["sectors"]),
                    rec["name"],
                    rec["state"],
                ]
            )
    points.sort(key=lambda p: -p[2])
    w.write_json(
        w.US / "ghgrp.json",
        {
            "fetched": w.now(),
            "year": GHGRP_YEAR,
            "source": f"EPA GHGRP {GHGRP_YEAR}: data summary (Direct Point Emitters, t CO2e, AR4 GWP) and Envirofacts PUB_DIM_FACILITY positions",
            "sectors": SECTORS,
            "points": points,
        },
    )
    print(
        f"  {len(rows):,} direct emitters, {len(points):,} placed; {sum(1 for r in out if r['position'] == 'Envirofacts'):,} positions from Envirofacts",
        flush=True,
    )
    tmp = TMP / "ghgrp_rows.json"
    tmp.write_text(json.dumps(out, default=str), encoding="utf-8")
    con = duckdb.connect()
    con.execute("INSTALL spatial; LOAD spatial")
    lk = w.lake_for(
        "ghgrp",
        args.lake,
        "EPA GHGRP data summary and Envirofacts",
        "public domain (US EPA)",
        "scripts/build_us_perf.py",
        str(GHGRP_YEAR),
    )
    lk.write(
        con,
        "facilities",
        f"""SELECT *, 'US' AS iso2, 'USA' AS iso3, state AS admin1,
                   CASE WHEN lon IS NOT NULL THEN ST_Point(lon::DOUBLE, lat::DOUBLE) END AS geometry
            FROM read_json('{tmp}', maximum_object_size = 1000000000, sample_size = -1)""",
        sort=["state", "facility_id"],
        partition={"year": str(GHGRP_YEAR)},
        geometry="geometry",
    )
    w.lake_done(lk, args.upload)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("what", choices=["perf", "ghgrp"])
    ap.add_argument("--lake", type=Path, default=Path("lake-out"))
    ap.add_argument("--upload", action="store_true")
    args = ap.parse_args()
    {"perf": perf, "ghgrp": ghgrp}[args.what](args)


if __name__ == "__main__":
    main()
