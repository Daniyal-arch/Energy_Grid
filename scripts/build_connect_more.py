"""Room to connect in six more countries (brief 21, group A): the operators' own figures.

    python scripts/build_connect_more.py fr|it|be|pt|ie|pl [--file F] [--lake lake-out] [--upload]

fr  Capareseau (RTE and the distribution operators, capareseau.fr; re-use allowed with the
    source named): every substation's S3REnR capacities as the site's region files publish them
    (`/medias/<id>` JSON per region: position, values by code) and the labels of the codes from
    its full CSV export. Passed through; numbers parsed where the value is a number.
it  Terna Econnextion via the dati.terna.it download center (JSON): requests to connect
    renewables (`FER`) and storage (`Accumuli`) by region and by province, MW and number of
    requests as published. Regions joined to NUTS-2 2024 (GISCO, simplified 0.01 degrees) by
    the table NUTS2 below (Trentino-Alto Adige = ITH1 + ITH2).
be  Elia, the capacity of the grid for new decentralised generation: elia.be answers scripts with
    HTTP 403, so the file comes from scripts/browser_fetch.py (`--file`). Columns as published.
pt  E-REDES `capacidade-rececao-rnd` (Opendatasoft, CC BY 4.0): reception capacity per HV/MV
    substation (MVA), as published; E-REDES gives no positions, so no map points.
ie  ESB Networks' capacity heatmap download (xlsx): every transformer's demand and generation
    capacity (MVA, MW) with its position, as published. The map keeps the 38 kV and higher
    stations; the lake keeps every row.
pl  PSE: the Art. 7(8l) list of connection requests, conditions and agreements on the
    transmission grid (xlsx, MW as published) and the list of PSE substations with no free
    connection bays (xlsx).

Each writes frontend/public/data/eu/connect/<cc>.json and a lake dataset (snapshot partition).
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

import wave5_common as w

OUT = w.EU / "connect"
TMP = Path("/tmp/wave6/connect")
TODAY = datetime.now(UTC).date().isoformat()

CAPARESEAU = "https://www.capareseau.fr"
AUTOCOMPLETE = "/medias/A2D4A43A-3FE4-E50F-27F8-4AC370326F94"
TERNA = "https://dati.terna.it/api/sitecore/dati/downloadcenter/records"
GISCO = "https://gisco-services.ec.europa.eu/distribution/v2/nuts/geojson/NUTS_RG_20M_2024_4326_LEVL_2.geojson"
EREDES = "https://e-redes.opendatasoft.com/api/explore/v2.1/catalog/datasets/capacidade-rececao-rnd/exports/json"
ESB_PAGE = (
    "https://www.esbnetworks.ie/our-services/tools---maps/generator-availability-capacity-map"
)
PSE_PAGE = "https://www.pse.pl/obszary-dzialalnosci/krajowy-system-elektroenergetyczny/wykaz-obiektow-planowanych-do-przylaczenia"

NUTS2 = {
    "PIEMONTE": ["ITC1"], "VALLE D'AOSTA": ["ITC2"], "LIGURIA": ["ITC3"], "LOMBARDIA": ["ITC4"],
    "TRENTINO-ALTO ADIGE": ["ITH1", "ITH2"], "VENETO": ["ITH3"], "FRIULI-VENEZIA GIULIA": ["ITH4"],
    "EMILIA-ROMAGNA": ["ITH5"], "TOSCANA": ["ITI1"], "UMBRIA": ["ITI2"], "MARCHE": ["ITI3"],
    "LAZIO": ["ITI4"], "ABRUZZO": ["ITF1"], "MOLISE": ["ITF2"], "CAMPANIA": ["ITF3"],
    "PUGLIA": ["ITF4"], "BASILICATA": ["ITF5"], "CALABRIA": ["ITF6"], "SICILIA": ["ITG1"],
    "SARDEGNA": ["ITG2"],
}  # fmt: skip


def number(v: object) -> float | None:
    """A published value as a number when it is one ("12.5", "1 234,5", 3); else None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, int | float):
        return float(v)
    t = str(v).strip().replace(" ", "").replace(" ", "")
    if re.fullmatch(r"-?\d+,\d+", t):
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def lake_rows(dataset: str, tables: list[tuple[str, list[dict], list[str], tuple[str, str] | None]],
              args: argparse.Namespace, *, source: str, licence: str) -> None:  # fmt: skip
    """Tables of rows (dicts) as snapshot partitions of lake/<dataset>/<table>:
    [(table, rows, sort, (lon, lat) columns or None)]."""
    import duckdb
    import lake

    TMP.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    con.execute("INSTALL spatial; LOAD spatial")
    lake.iso_codes(con)
    lk = w.lake_for(dataset, args.lake, source, licence, "scripts/build_connect_more.py", TODAY)
    for table, rows, sort, geometry in tables:
        tmp = TMP / f"{dataset}_{table}.json"
        tmp.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
        geo = ""
        if geometry:
            lon, lat = geometry
            geo = f", CASE WHEN {lon} IS NOT NULL THEN ST_Point({lon}, {lat}) END AS geometry"
        lk.write(
            con, table,
            f"SELECT *{geo} FROM read_json('{tmp}', sample_size = -1, maximum_object_size = 1000000000)",
            sort=sort, partition={"snapshot": TODAY}, geometry="geometry" if geometry else None,
        )  # fmt: skip
    w.lake_done(lk, args.upload)


# ---------------------------------------------------------------- France


def capareseau_labels(csv_text: str) -> dict[str, str]:
    """The CSV export's first line holds the codes, the second their labels."""
    lines = csv_text.lstrip("﻿").splitlines()
    split = lambda s: [c.strip().strip('"') for c in re.split(r';(?=(?:[^"]*"[^"]*")*[^"]*$)', s)]  # noqa: E731
    return dict(zip(split(lines[0]), split(lines[1]), strict=False))


def capareseau_site(s: dict) -> dict:
    """One substation of a region file: identity, position, and every value as published."""
    vals = {k: v for k, v in (s.get("values") or {}).items() if v not in ("", None)}
    return {
        "code": s.get("code"),
        "name": s.get("name"),
        "region": s.get("territory_name"),
        "lon": s.get("X"),
        "lat": s.get("Y"),
        "u_max_kv": s.get("u_max"),
        "updated": s.get("updated"),
        "grd1": (s.get("grd1") or {}).get("name"),
        "grd2": (s.get("grd2") or {}).get("name"),
        "values": vals,
    }


def build_fr(args: argparse.Namespace) -> None:
    names = w.get_json(CAPARESEAU + AUTOCOMPLETE)
    regions = sorted({x["territory_code"] for x in names if x.get("territory_code")})  # type: ignore[union-attr]
    print(f"  {len(names)} names, regions {regions}", flush=True)  # type: ignore[arg-type]
    sites: list[dict] = []
    labels: dict[str, str] = {}
    csv_url = ""
    import httpx

    for code in regions:
        html = httpx.get(f"{CAPARESEAU}/region/{code}", headers=w.UA, timeout=120).text
        for media in dict.fromkeys(re.findall(r"/medias/([A-F0-9-]{36})", html)):
            if CAPARESEAU + "/medias/" + media == CAPARESEAU + AUTOCOMPLETE:
                continue
            r = httpx.get(f"{CAPARESEAU}/medias/{media}", headers=w.UA, timeout=120)
            disp = r.headers.get("content-disposition", "")
            if disp.endswith(".csv") or ".csv" in disp:
                if not labels:
                    labels = capareseau_labels(r.content.decode("utf-8-sig", "replace"))
                    csv_url = f"{CAPARESEAU}/medias/{media} ({disp.split('filename=')[-1]})"
                continue
            try:
                data = r.json()
            except ValueError:
                continue
            if isinstance(data, list) and data and "X" in data[0]:
                sites += [capareseau_site(s) for s in data]
                print(f"  region {code}: {len(data)} substations", flush=True)
    by_code = {s["code"]: s for s in sites}
    sites = sorted(by_code.values(), key=lambda s: s["code"] or "")
    if len(sites) < 1000:
        raise SystemExit(f"only {len(sites)} substations")
    codes = sorted({k for s in sites for k in s["values"]})
    payload = {
        "source": "Capareseau (RTE and the distribution grid operators), capareseau.fr",
        "url": csv_url,
        "codes": codes,
        "labels": {c: labels.get(c, c) for c in codes},
        # [code, name, region, lon, lat, u max kV, updated, values per code (number, text or null)]
        "sites": [
            [s["code"], s["name"], s["region"], s["lon"], s["lat"], s["u_max_kv"], s["updated"],
             [(number(s["values"][c]) if number(s["values"][c]) is not None else s["values"][c]) if c in s["values"] else None for c in codes]]
            for s in sites
        ],
    }  # fmt: skip
    w.write_json(OUT / "fr.json", {"fetched": w.now(), **payload})
    rows = [
        {**{k: v for k, v in s.items() if k != "values"}, "iso2": "FR", "iso3": "FRA", "admin1": s["region"],
         **{c.lower(): (str(s["values"][c]) if c in s["values"] else None) for c in codes}}
        for s in sites
    ]  # fmt: skip
    lake_rows("capareseau", [("substations", rows, ["code"], ("lon", "lat"))], args, source="capareseau.fr (RTE, Enedis and the local DSOs)",
              licence="re-use allowed with the source named (capareseau.fr legal notice)")  # fmt: skip


# ---------------------------------------------------------------- Italy


def terna(dataset: str, view: str) -> tuple[list[str], list[list]]:
    q = f"{TERNA}?filterDataset={dataset}&filterViewBy={view}&orderByColumn=Potenza%20(MW)&orderByDir=desc&pageIndex=0&pageSize=5000"
    d = w.get_json(q, timeout=180)
    return list(d["Columns"]), d["Data"]  # type: ignore[index]


def it_regions(rows: list[list], cols: list[str]) -> dict[str, dict[str, list[float]]]:
    """{region: {source: [MW, requests]}} from the per-region answer (requests may be absent)."""
    out: dict[str, dict[str, list[float]]] = defaultdict(dict)
    ri, si, mi = cols.index("Regione"), cols.index("Fonte"), cols.index("Potenza (MW)")
    ni = cols.index("Numero Pratiche") if "Numero Pratiche" in cols else None
    for r in rows:
        out[r[ri]][r[si]] = [round(r[mi], 3), r[ni] if ni is not None else None]
    return dict(out)


def build_it(args: argparse.Namespace) -> None:
    import shapely
    from shapely.geometry import shape

    out: dict[str, object] = {
        "source": "Terna, Econnextion (dati.terna.it download center)",
        "url": TERNA,
    }
    lake: list[dict] = []
    for ds, label in (("FER", "renewables"), ("Accumuli", "storage")):
        cols, rows = terna(ds, "Region")
        pcols, prows = terna(ds, "Province")
        print(
            f"  {ds}: {len(rows)} region rows, {len(prows)} province rows; columns {cols}",
            flush=True,
        )
        out[label] = {
            "regions": it_regions(rows, cols),
            "provinces": [dict(zip(pcols, r, strict=False)) for r in prows],
        }
        lake += [
            {"dataset": ds, "view": "region", **dict(zip(cols, r, strict=False))} for r in rows
        ]
        lake += [
            {"dataset": ds, "view": "province", **dict(zip(pcols, r, strict=False))} for r in prows
        ]
    gj = w.get_json(GISCO, timeout=300)
    shapes: dict[str, list] = {}
    for f in gj["features"]:  # type: ignore[index]
        p = f["properties"]
        if p.get("CNTR_CODE") == "IT":
            shapes[p["NUTS_ID"]] = w.rings(shapely.simplify(shape(f["geometry"]), 0.01), 2)
    out["nuts2"] = {reg: ids for reg, ids in NUTS2.items()}
    out["shapes"] = shapes
    missing = [r for r in out["renewables"]["regions"] if r not in NUTS2]  # type: ignore[index]
    if missing:
        raise SystemExit(f"regions without NUTS-2: {missing}")
    w.write_json(OUT / "it.json", {"fetched": w.now(), **out})
    rows = [
        {"dataset": r["dataset"], "view": r["view"], "region": r.get("Regione"), "province": r.get("Provincia"),
         "source": r.get("Fonte"), "plant_type": r.get("Tipo Impianto"), "mw": r.get("Potenza (MW)"),
         "requests": r.get("Numero Pratiche"), "iso2": "IT", "iso3": "ITA", "admin1": r.get("Regione")}
        for r in lake
    ]  # fmt: skip
    lake_rows("econnextion", [("requests", rows, ["dataset", "view", "region", "province", "source"], None)], args,
              source="Terna Econnextion (dati.terna.it)", licence="Terna terms (not found on dati.terna.it)")  # fmt: skip


# ---------------------------------------------------------------- Belgium


def build_be(args: argparse.Namespace) -> None:
    if not args.file:
        raise SystemExit(
            "be: the Elia file comes from scripts/browser_fetch.py elia_capacity (--file)"
        )
    import openpyxl

    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    sheets = {}
    lake: list[dict] = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
        head_i = next((i for i, r in enumerate(rows) if sum(isinstance(v, str) for v in r) >= 3), 0)
        head = (
            [str(v).strip() if v is not None else f"col{j}" for j, v in enumerate(rows[head_i])]
            if rows
            else []
        )
        body = [
            [v if isinstance(v, int | float | str) or v is None else str(v) for v in r]
            for r in rows[head_i + 1 :]
        ]
        sheets[ws.title] = {"columns": head, "rows": body}
        lake += [
            {
                "sheet": ws.title,
                **{h: (str(v) if v is not None else None) for h, v in zip(head, r, strict=False)},
            }
            for r in body
        ]
        print(f"  sheet {ws.title}: {len(body)} rows, columns {head[:12]}", flush=True)
    w.write_json(
        OUT / "be.json",
        {
            "fetched": w.now(),
            "source": "Elia, capacity for decentralised generation",
            "file": Path(args.file).name,
            "sheets": sheets,
        },
    )
    lake_rows(
        "elia_hosting",
        [("capacity", lake, ["sheet"], None)],
        args,
        source="Elia (elia.be)",
        licence="Elia terms of use",
    )


# ---------------------------------------------------------------- Portugal

PT_KEEP = [
    "potencia_de_ligacao_ligado_mva_ultimo_trimestre",
    "potencia_de_ligacao_comprometido_mva_ultimo_trimestre",
    "potencia_de_ligacao_em_confirmacao_mva_ultimo_trimestre",
    "capacidade_de_recepcao_at_mva_ultimo_trimestre",
    "capacidade_de_recepcao_mt_at_mva_ultimo_trimestre",
    "capacidade_de_recepcao_mt_30kv_mva_ultimo_trimestre",
    "capacidade_de_recepcao_mt_15kv_mva_ultimo_trimestre",
    "capacidade_de_recepcao_mt_10kv_mva_ultimo_trimestre",
]


def pt_row(r: dict) -> list:
    """[code, name, type, district, municipality, NUTS-3 name, RNT bus, values of PT_KEEP]."""
    return [r.get("codigo"), r.get("instalacao"), r.get("tipo_de_instalacao"), r.get("distrito"), r.get("concelho"),
            r.get("nut3"), r.get("ligacao_rnt_barramento_60kv_ultimo_trimestre") or r.get("ligacao_rnt_barramento_60kv_rari"),
            [number(r.get(k)) for k in PT_KEEP]]  # fmt: skip


def build_pt(args: argparse.Namespace) -> None:
    rows = w.get_json(EREDES, timeout=180)
    assert isinstance(rows, list)
    dates = Counter(r.get("data_ultimo_trimestre") for r in rows)
    print(f"  {len(rows)} substations; quarter dates {dates.most_common(3)}", flush=True)
    w.write_json(OUT / "pt.json", {
        "fetched": w.now(), "source": "E-REDES open data, capacidade-rececao-rnd (CC BY 4.0)",
        "date": dates.most_common(1)[0][0], "columns": PT_KEEP, "substations": [pt_row(r) for r in rows],
    })  # fmt: skip
    lake_rows("eredes_capacity", [("substations", [{**r, "iso2": "PT", "iso3": "PRT", "admin1": r.get("distrito")} for r in rows], ["codigo"], None)], args,
              source="E-REDES open data", licence="CC BY 4.0")  # fmt: skip


# ---------------------------------------------------------------- Ireland

ESB_COLS = {
    "Station Name": "name", "Transformer GroupID": "group", "Primary kV": "primary_kv",
    "Secondary voltage(s)": "secondary", "Voltage Class": "voltage_class",
    "Transformer Configuration": "configuration", "Installed Capacity MVA": "installed_mva",
    "Demand FirmCapacity MVA": "demand_firm_mva", "Demand Available MVA": "demand_available_mva",
    "Parent Available MVA": "parent_available_mva", "Demand Parent Constraint": "demand_parent_constraint",
    "Generation Firm Capacity MW": "gen_firm_mw", "Generation NonFirm Capacity MW": "gen_nonfirm_mw",
    "Generation Total Committed MW": "gen_committed_mw", "Gen Available Firm MW": "gen_available_firm_mw",
    "Gen Available NonFirm MW": "gen_available_nonfirm_mw", "Generation Parent Constraint": "gen_parent_constraint",
    "Parent Feeder": "parent_feeder", "Parent Station": "parent_station",
    "TSO Interface Station": "tso_station", "Comment": "comment", "Latitude": "lat", "Longitude": "lon",
}  # fmt: skip


def esb_rows(rows: list[tuple]) -> list[dict]:
    """The heatmap sheet: the header row holds 'Station Name'; values as published."""
    hi = next(i for i, r in enumerate(rows) if r and "Station Name" in r)
    head = [str(v).strip() if v is not None else "" for v in rows[hi]]
    out = []
    for r in rows[hi + 1 :]:
        rec = {ESB_COLS[h]: v for h, v in zip(head, r, strict=False) if h in ESB_COLS}
        if not rec.get("name"):
            continue
        for k in list(rec):
            if k.endswith(("_mva", "_mw")) or k in ("lat", "lon"):
                rec[k] = number(rec[k])
            elif rec[k] is not None:
                rec[k] = str(rec[k]).strip()
        out.append(rec)
    return out


def build_ie(args: argparse.Namespace) -> None:
    import httpx
    import openpyxl

    html = httpx.get(ESB_PAGE, headers=w.UA, timeout=120, follow_redirects=True).text
    hits = re.findall(r"""["'](https://[^"']*heatmap[^"']*\.xlsx[^"']*)["']""", html, re.I)
    if not hits:
        raise SystemExit("ESB Networks: no heatmap xlsx link on the page")
    url = hits[0].replace("&amp;", "&")
    path = w.download(url, TMP / "esb_heatmap.xlsx")
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Heatmap data"]
    rows = esb_rows(list(ws.iter_rows(values_only=True)))
    classes = Counter(r.get("voltage_class") for r in rows)
    print(f"  {len(rows)} transformers; voltage classes {classes.most_common()}", flush=True)
    keep = [r for r in rows if r.get("voltage_class") != "LV" and r.get("lon") is not None]
    cols = [
        "installed_mva",
        "demand_available_mva",
        "gen_available_firm_mw",
        "gen_available_nonfirm_mw",
        "gen_committed_mw",
    ]
    w.write_json(OUT / "ie.json", {
        "fetched": w.now(), "source": "ESB Networks, Available Capacity Heatmap (download)", "url": url,
        "file": Path(url.split("?")[0]).name, "transformers_total": len(rows), "classes": dict(classes), "columns": cols,
        # [name, primary kV, voltage class, lon, lat, values of columns, comment]
        "stations": [[r["name"], r.get("primary_kv"), r.get("voltage_class"), round(r["lon"], 5), round(r["lat"], 5),
                      [r.get(c) for c in cols], r.get("comment") or None] for r in keep],
    })  # fmt: skip
    lake_rows("esb_heatmap", [("transformers", [{**r, "iso2": "IE", "iso3": "IRL", "admin1": None} for r in rows], ["name"], ("lon", "lat"))], args,
              source="ESB Networks capacity heatmap", licence="ESB Networks website terms")  # fmt: skip


# ---------------------------------------------------------------- Poland


def pse_requests(rows: list[tuple]) -> list[dict]:
    """The Art. 7(8l) sheet: header row has 'ID Obiektu'; MW columns and dates as published."""
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "ID Obiektu")
    keys = ["id", "name", "operator", "applicant", "location", "voltage_kv", "multiple", "type",
            "mw_injected", "mw_withdrawn", "pv_mw", "wind_mw", "storage_discharge_mw", "storage_charge_mw",
            "demand_mw", "other_mw", "status", "applied", "advance_paid", "conditions_issued", "refused",
            "refusal_reason", "conditions_expired", "agreement", "supply_start", "agreement_ended"]  # fmt: skip
    out = []
    for r in rows[hi + 1 :]:
        if not r or not r[1]:
            continue
        rec = {}
        for k, v in zip(keys, r, strict=False):
            if isinstance(v, datetime):
                v = v.date().isoformat()
            if isinstance(v, str):
                v = re.sub(r"\s+", " ", v).strip()
                v = None if v in ("-", "") else v
            rec[k] = v
        out.append(rec)
    return out


def pse_full(rows: list[tuple]) -> list[dict]:
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "ZKO PSE")
    return [
        {"area": r[0], "symbol": r[1], "station": r[2], "switchgear": r[3], "free_bays": r[4]}
        for r in rows[hi + 1 :]
        if r and r[2]
    ]


def build_pl(args: argparse.Namespace) -> None:
    import httpx
    import openpyxl

    html = httpx.get(PSE_PAGE, headers=w.UA, timeout=120, follow_redirects=True).text
    links = list(
        dict.fromkeys(
            h.replace("&amp;", "&") for h in re.findall(r'href="([^"]+\.xlsx[^"]*)"', html)
        )
    )
    art = sorted(h for h in links if "Art_7_ust_8l" in h)
    full = sorted(h for h in links if "Lista_rozdzielni" in h or "Lista+rozdzielni" in h)
    print(f"  Art. 7(8l) files {art[-2:]}; substation lists {full[-2:]}", flush=True)
    if not art or not full:
        raise SystemExit("PSE: files not found on the page")
    absu = lambda h: h if h.startswith("http") else "https://www.pse.pl" + h  # noqa: E731
    a = w.download(absu(art[-1]), TMP / "pse_art7.xlsx")
    b = w.download(absu(full[-1]), TMP / "pse_full.xlsx")
    reqs = pse_requests(
        list(
            openpyxl.load_workbook(a, read_only=True, data_only=True)
            .worksheets[0]
            .iter_rows(values_only=True)
        )
    )
    subs = pse_full(
        list(
            openpyxl.load_workbook(b, read_only=True, data_only=True)
            .worksheets[0]
            .iter_rows(values_only=True)
        )
    )
    as_of = re.search(r"stan_na_(\d\d)_(\d\d)_(\d{4})", art[-1])
    date = f"{as_of.group(3)}-{as_of.group(2)}-{as_of.group(1)}" if as_of else None
    print(
        f"  {len(reqs)} requests (as of {date}); {len(subs)} switchgears without free bays",
        flush=True,
    )
    num = lambda v: v if isinstance(v, int | float) else None  # noqa: E731 - "129; 240" stays text
    w.write_json(OUT / "pl.json", {
        "fetched": w.now(), "source": "PSE S.A., information under Art. 7(8l) of the Energy Law; list of switchgears without free bays",
        "date": date, "url": absu(art[-1]), "url_full": absu(full[-1]),
        # [name, type, status, voltage, MW injected, MW withdrawn, location]
        "requests": [[r["name"], r["type"], r["status"], r["voltage_kv"] if isinstance(r["voltage_kv"], int | float) else (str(r["voltage_kv"]) if r["voltage_kv"] else None),
                      num(r["mw_injected"]), num(r["mw_withdrawn"]), r["location"]] for r in reqs],
        "full": [[s["station"], s["symbol"], s["switchgear"], s["area"]] for s in subs],
    })  # fmt: skip
    lake_rows("pse_capacity", [
        ("requests", [{**{k: (str(v) if v is not None else None) for k, v in r.items()}, "iso2": "PL", "iso3": "POL", "admin1": None} for r in reqs], ["name"], None),
        ("full_substations", [{**s, "iso2": "PL", "iso3": "POL", "admin1": None} for s in subs], ["station"], None),
    ], args, source="PSE S.A. (pse.pl)", licence="PSE terms (public information under Art. 7(8l) Energy Law)")  # fmt: skip


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("country", choices=["fr", "it", "be", "pt", "ie", "pl"])
    ap.add_argument("--file", type=Path)
    ap.add_argument("--lake", type=Path, default=Path("lake-out"))
    ap.add_argument("--upload", action="store_true")
    args = ap.parse_args()
    {
        "fr": build_fr,
        "it": build_it,
        "be": build_be,
        "pt": build_pt,
        "ie": build_ie,
        "pl": build_pl,
    }[args.country](args)


if __name__ == "__main__":
    main()
