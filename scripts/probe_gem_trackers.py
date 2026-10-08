"""Probe the GEM tracker files in data/world/gem/ (scripts/fetch_gem.py): sheets, columns,
row counts and coordinates of each Excel file, and what the GIS zips hold."""

from __future__ import annotations

import sys
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1] / "data" / "world" / "gem"
SKIP = {"integrated-power", "energy-ownership"}  # known / large: probed separately


def excel(path: Path) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        cols = [str(c) for c in header if c is not None]
        sample = []
        n = 0
        status = Counter()
        st_col = next(
            (i for i, c in enumerate(header) if c and str(c).strip().lower() == "status"), None
        )
        for row in rows:
            n += 1
            if len(sample) < 1:
                sample.append(row)
            if st_col is not None and st_col < len(row):
                status[str(row[st_col])] += 1
        print(f"  sheet {ws.title!r}: {n} rows, {len(cols)} columns")
        print(f"    columns: {cols[:60]}")
        if sample:
            pairs = [
                (str(h), v) for h, v in zip(header, sample[0], strict=False) if v not in (None, "")
            ]
            print(f"    first row: {str(pairs[:28])[:900]}")
        if status:
            print(f"    status: {status.most_common(10)}")
    wb.close()


def main() -> None:
    only = set(sys.argv[1:])
    for folder in sorted(ROOT.iterdir()):
        if not folder.is_dir() or folder.name in SKIP or (only and folder.name not in only):
            continue
        print(f"=== {folder.name}")
        for f in sorted(folder.iterdir()):
            print(f"- {f.name} ({f.stat().st_size / 1e6:.1f} MB)")
            if f.suffix == ".xlsx":
                excel(f)
            elif f.suffix == ".zip":
                with zipfile.ZipFile(f) as z:
                    for info in z.infolist()[:12]:
                        print(f"    {info.filename} ({info.file_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
